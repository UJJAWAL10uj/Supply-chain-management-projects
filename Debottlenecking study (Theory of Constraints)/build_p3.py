"""
Project 3: Theory-of-Constraints debottlenecking study (Excel model + SimPy results).
Deterministic capacity, constraint, scenario and economics logic are Excel formulas;
stochastic results are pasted from simulate.py / run_experiments.py (results.json).
"""
import json

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

import simulate as S

OUT = "/mnt/user-data/outputs/Debottlenecking_Study_Enclosure_Line.xlsx"
P1 = "/mnt/user-data/outputs/Apple_Sourcing_TCO_Model_Aluminium_Enclosure.xlsx"
RES = json.load(open("/home/claude/results.json"))

# ---- pull price and material from Project 1 model (base India column)
w1 = load_workbook(P1, data_only=True)["Calc"]
col = next(c for c in range(3, 40) if w1.cell(6, c).value == "Base case|India")
P1_PRICE = round(w1.cell(51, col).value, 3)   # ex-works price before incentive
P1_MATERIAL = round(w1.cell(41, col).value, 3)  # net material

FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10)
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
WHITE_B = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=10, italic=True, color="595959")
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")
TOT = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="top")
INT = '#,##0;(#,##0);"-"'
N1 = '#,##0.0;(#,##0.0);"-"'
N2 = '#,##0.00;(#,##0.00);"-"'
N3 = '0.000;(0.000);"-"'
PCT = '0.0%;(0.0%);"-"'
USD0 = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
amb = PatternFill("solid", bgColor="FFE699", fgColor="FFE699")
grn = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")

wb = Workbook()


def put(ws, ref, v, font=BLACK, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def hdr(ws, row, labels, c0=1):
    for i, t in enumerate(labels):
        put(ws, f"{L(c0 + i)}{row}", t, WHITE_B, fill=HDR, align=CENTER, border=BOX)


def sec(ws, row, text, n):
    for i in range(1, n + 1):
        ws[f"{L(i)}{row}"].fill = SEC
    put(ws, f"A{row}", text, BOLD, fill=SEC)


def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


# ================================================================ README
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 26, "C": 112})
put(ws, "B2", "Debottlenecking Study (Theory of Constraints): Aluminium Enclosure Production Line", TITLE)
put(ws, "B3", "SIMULATED / ILLUSTRATIVE DATA. Not a real plant. Stage parameters are benchmark-informed assumptions; replace with time-study data.", SUB)
rows = [
    ("Question", "One production module must deliver 10,000 good enclosures a week. Where is the constraint, how much output does it cost, which fixes recover it, and what do they cost?"),
    ("Line", "Blank saw > CNC OP10 rough > CNC OP20 finish > deburr / polish > anodising (48-unit rack batches) > inspection > packing. One module is 10% of the 5-million-unit-a-year line modelled in Project 1."),
    ("Two models, one answer", "Excel (this workbook) does the rate-based Theory of Constraints analysis: capacity, utilisation, takt, OEE, constraint, scenario output and economics, all as formulas. A SimPy discrete-event simulation (simulate.py, 8 replications per scenario) adds what averages hide: breakdowns, cycle-time variability, queues, WIP and lead time. Simulation outputs are pasted on the SimResults sheet."),
    ("Five focusing steps", "1 Identify the constraint (BaseAnalysis). 2 Exploit it: get more from what exists (SMED, TPM, yield at the constraint). 3 Subordinate everything else (drum-buffer-rope release). 4 Elevate it: add capacity only if still needed (third line, 7-day working). 5 Repeat: the constraint moves (Fixes sheet shows the next one)."),
    ("Sheets", "Dashboard > Inputs > BaseAnalysis > Fixes (deterministic scenarios) > SimResults (SimPy output, rope sweep) > Economics (gain vs cost) > Sources."),
    ("How to use", "Change only BLUE cells (stage parameters, fix parameters, costs, 0/1 scenario flags). Everything except SimResults recalculates. If you change stage parameters, re-run simulate.py and run_experiments.py and repaste SimResults so the simulation matches."),
    ("Colour code", "Blue = input. Black = formula. Green = link to another sheet. Yellow fill = key assumption."),
    ("Cost-benefit logic", "Gain = extra sellable good units per week (capped at demand) x 50 weeks x contribution per unit. Contribution = Project 1 ex-works price - net material - assumed variable conversion cost. Costs are illustrative one-time and annual amounts; one-time costs are spread over 3 years. Gains assume the customer will buy the recovered units."),
    ("Limitations", "Simulated line, not a real one. Lots of 6 units move between stages. Failures are operating-time based. Costs of fixes are assumptions to be replaced by quotes. Weekly gain is demand-capped, so fixes beyond the first are worth little unless demand grows."),
]
r = 5
for k, v in rows:
    put(ws, f"B{r}", k, BOLD, align=WRAP)
    put(ws, f"C{r}", v, BLACK, align=WRAP)
    ws.row_dimensions[r].height = 62 if len(v) > 230 else 46
    r += 1

# ================================================================ INPUTS
wi = wb.create_sheet("Inputs")
wi.sheet_view.showGridLines = False
widths(wi, {"A": 46, "B": 30, "C": 14, "D": 14, "E": 16, "F": 14, "G": 14, "H": 13, "I": 11, "J": 11, "K": 70})
put(wi, "A1", "Inputs and assumptions (simulated line)", TITLE)
put(wi, "A2", "Blue = input. Yellow = key assumption. Stage values are illustrative, benchmark-informed (see Sources).", SUB)
sec(wi, 4, "Global parameters", 11)
glob = [
    (5, "Planned production hours per week", 96, "h/week", INT, "2 shifts x 8 h x 6 days", False),
    (6, "Planned production minutes per week", "=B5*60", "min/week", INT, "Formula", False),
    (7, "Weekly customer demand (good units)", S.DEMAND, "units/week", INT, "Module target: 10% of Project 1's 5M units a year", True),
    (8, "Production weeks per year", 50, "weeks", INT, "Assumption", False),
    (9, "Anodising rack size", S.RACK, "units", INT, "Assumption: units per rack batch", False),
    (10, "Ex-works price per unit", P1_PRICE, "USD", USD2, "From Project 1 model: India base case, before ECMS incentive", False),
    (11, "Net material cost per unit", P1_MATERIAL, "USD", USD2, "From Project 1 model: India base case, net of chip-scrap credit", False),
    (12, "Variable conversion cost per unit", 0.60, "USD", USD2, "Assumption: consumables, chemicals, energy that vary with output", True),
    (13, "Contribution per extra good unit", "=B10-B11-B12", "USD", USD2, "Formula: price - material - variable conversion", False),
    (14, "Amortisation period for one-time costs", 3, "years", INT, "Assumption", False),
    (15, "Scale factor, module to full 5M-unit line", 10, "x", INT, "Assumption: ten identical modules", False),
]
for row, lab, val, unit, fmt, note, key in glob:
    put(wi, f"A{row}", lab, BLACK, border=BOX)
    is_formula = isinstance(val, str) and val.startswith("=")
    put(wi, f"B{row}", val, BLACK if is_formula else BLUE, fmt, KEY if key else None, border=BOX, align=RIGHT)
    put(wi, f"C{row}", unit, BLACK, border=BOX)
    put(wi, f"K{row}", note, BLACK, align=WRAP)

sec(wi, 17, "Stage parameters", 11)
hdr(wi, 18, ["#", "Stage", "Type", "Machines / lines", "Ideal cycle per unit (min)", "Rack batch time (min)", "Availability", "Performance", "Yield", "MTTR (min)", "Basis"])
wi.row_dimensions[18].height = 44
basis = [
    "Fast, one saw; not a constraint",
    "Rough milling, 7 machines",
    "Finish milling, 9 machines; availability 90% = world-class threshold [W1]",
    "Semi-automatic deburr / polish stations",
    "Batch process: 2 rack lines, 52 min per 48-unit rack incl. load, pretreat, anodise, dye, seal",
    "100% visual plus sampled CMM",
    "Pack and label",
]
for i, st in enumerate(S.BASE_STAGES):
    r_ = 19 + i
    put(wi, f"A{r_}", i + 1, BLACK, INT, align=CENTER, border=BOX)
    put(wi, f"B{r_}", st["name"], BOLD, border=BOX)
    put(wi, f"C{r_}", "Batch" if st["kind"] == "batch" else "Per unit", BLACK, border=BOX)
    put(wi, f"D{r_}", st["n"], BLUE, INT, KEY, CENTER, BOX)
    if st["kind"] == "batch":
        put(wi, f"E{r_}", f"=F{r_}/$B$9", BLACK, N3, align=CENTER, border=BOX)
        put(wi, f"F{r_}", st["batch_min"], BLUE, N1, KEY, CENTER, BOX)
    else:
        put(wi, f"E{r_}", st["cycle"], BLUE, N2, KEY, CENTER, BOX)
        put(wi, f"F{r_}", None, BLACK, border=BOX)
    put(wi, f"G{r_}", st["A"], BLUE, PCT, KEY, CENTER, BOX)
    put(wi, f"H{r_}", st["P"], BLUE, PCT, KEY, CENTER, BOX)
    put(wi, f"I{r_}", st["Y"], BLUE, PCT, KEY, CENTER, BOX)
    put(wi, f"J{r_}", st["mttr"], BLUE, INT, align=CENTER, border=BOX)
    put(wi, f"K{r_}", basis[i], BLACK, align=WRAP)

sec(wi, 27, "Fix parameters and illustrative costs (all cost values are assumptions)", 11)
hdr(wi, 28, ["Code", "Fix", "Parameter", "Unit", "One-time cost (USD)", "Annual cost (USD)", "", "", "", "", "What it does (TOC step)"])
wi.row_dimensions[28].height = 36
fixes = [
    ("DBR", "Drum-buffer-rope release", RES["rope_chosen"], "units ahead of constraint", 10000, 0, f"Subordinate: release only when the buffer ahead of anodising is below the rope; cuts WIP and lead time. Rope {RES['rope_chosen']} units at base capacity, re-sized to {RES['rope_resized']} after the constraint is sped up"),
    ("YIELD", "Anodising yield improvement", S.FIX_PARAMS["yield_new"], "anodising yield", 20000, 60000, "Exploit: scrap at the constraint wastes constraint hours; better chemistry and process control"),
    ("SMED", "SMED on rack load / unload", S.FIX_PARAMS["smed_batch_min"], "min per rack batch", 40000, 0, "Exploit: faster rack changeover raises constraint capacity at low cost"),
    ("SEVEN", "Anodising 7 days a week", "=7/6", "capacity multiplier", 0, 120000, "Elevate: extra constraint hours via weekend working"),
    ("LINE3", "Add third anodising line", S.FIX_PARAMS["line3_lines"], "anodising lines", 450000, 90000, "Elevate: parallel station, capex"),
    ("TPM", "TPM at anodising", S.FIX_PARAMS["tpm_A"], "anodising availability", 10000, 35000, "Exploit: preventive maintenance cuts breakdown time at the constraint"),
]
for i, f in enumerate(fixes):
    r_ = 29 + i
    put(wi, f"A{r_}", f[0], BOLD, border=BOX)
    put(wi, f"B{r_}", f[1], BLACK, border=BOX)
    is_f = isinstance(f[2], str)
    fm = PCT if f[0] in ("YIELD", "TPM") else (N3 if f[0] == "SEVEN" else INT)
    put(wi, f"C{r_}", f[2], BLACK if is_f else BLUE, fm, KEY, CENTER, BOX)
    put(wi, f"D{r_}", f[3], BLACK, border=BOX)
    put(wi, f"E{r_}", f[4], BLUE, USD0, KEY, border=BOX)
    put(wi, f"F{r_}", f[5], BLUE, USD0, KEY, border=BOX)
    put(wi, f"K{r_}", f[6], BLACK, align=WRAP)
    wi.row_dimensions[r_].height = 30
wi.freeze_panes = "A4"

# ================================================================ BASE ANALYSIS
ba = wb.create_sheet("BaseAnalysis")
ba.sheet_view.showGridLines = False
widths(ba, {"A": 6, "B": 28, "C": 11, "D": 13, "E": 14, "F": 14, "G": 13, "H": 15, "I": 14, "J": 10, "K": 15, "L": 14, "M": 11})
put(ba, "A1", "Step 1: identify the constraint (rate-based analysis at target demand)", TITLE)
put(ba, "A2", "Capacity = machines x planned minutes x availability x performance / ideal cycle. Utilisation above 100% means the stage cannot keep up with demand.", SUB)
hdr(ba, 4, ["#", "Stage", "Machines / lines", "Cumulative yield before stage", "Load at target (units/wk)", "Capacity (units/wk)", "Utilisation at target", "Effective cycle per unit (min)", "Stage takt (min/unit)", "OEE", "Starts capacity (starts/wk)", "Constraint?", "Util. rank"])
ba.row_dimensions[4].height = 58
for i in range(7):
    r_, ir = 5 + i, 19 + i
    put(ba, f"A{r_}", f"=Inputs!A{ir}", GREEN, INT, align=CENTER, border=BOX)
    put(ba, f"B{r_}", f"=Inputs!B{ir}", GREEN, border=BOX)
    put(ba, f"C{r_}", f"=Inputs!D{ir}", GREEN, INT, align=CENTER, border=BOX)
    put(ba, f"D{r_}", 1 if i == 0 else f"=D{r_ - 1}*Inputs!I{ir - 1}", BLACK, N3, align=CENTER, border=BOX)
    put(ba, f"E{r_}", f"=$D$15*D{r_}", BLACK, INT, border=BOX)
    put(ba, f"F{r_}", f"=C{r_}*Inputs!$B$6*Inputs!G{ir}*Inputs!H{ir}/Inputs!E{ir}", BLACK, INT, border=BOX)
    put(ba, f"G{r_}", f"=E{r_}/F{r_}", BOLD, PCT, align=CENTER, border=BOX)
    put(ba, f"H{r_}", f"=Inputs!E{ir}/(C{r_}*Inputs!G{ir}*Inputs!H{ir})", BLACK, N3, align=CENTER, border=BOX)
    put(ba, f"I{r_}", f"=Inputs!$B$6/E{r_}", BLACK, N3, align=CENTER, border=BOX)
    put(ba, f"J{r_}", f"=Inputs!G{ir}*Inputs!H{ir}*Inputs!I{ir}", BLACK, PCT, align=CENTER, border=BOX)
    put(ba, f"K{r_}", f"=F{r_}/D{r_}", BLACK, INT, border=BOX)
    put(ba, f"L{r_}", f'=IF(K{r_}=MIN($K$5:$K$11),"CONSTRAINT","")', BOLD, align=CENTER, border=BOX)
    put(ba, f"M{r_}", f"=RANK(G{r_},$G$5:$G$11,0)", BLACK, INT, align=CENTER, border=BOX)
ba.conditional_formatting.add("G5:G11", FormulaRule(formula=["G5>=1"], fill=red))
ba.conditional_formatting.add("G5:G11", FormulaRule(formula=["G5>=0.9"], fill=amb))
ba.conditional_formatting.add("G5:G11", FormulaRule(formula=["G5<0.9"], fill=grn))
ba.conditional_formatting.add("L5:L11", FormulaRule(formula=['L5="CONSTRAINT"'], fill=red))
sec(ba, 13, "Line summary", 13)
summ = [
    (14, "Total line yield (first-pass)", "=PRODUCT(Inputs!I19:I25)", PCT),
    (15, "Starts per week needed to meet demand", "=Inputs!B7/D14", INT),
    (16, "Takt time per good unit (min)", "=Inputs!B6/Inputs!B7", N3),
    (17, "Takt time per start (min)", "=Inputs!B6/D15", N3),
    (18, "Maximum starts per week the line can support", "=MIN(K5:K11)", INT),
    (19, "Constraint stage", "=INDEX(B5:B11,MATCH(D18,K5:K11,0))", None),
    (20, "Maximum good output per week", "=D18*D14", INT),
    (21, "Shortfall versus demand (units/wk)", "=D20-Inputs!B7", INT),
    (22, "Output as % of demand", "=D20/Inputs!B7", PCT),
    (23, "OEE of the constraint stage", "=INDEX(J5:J11,MATCH(D18,K5:K11,0))", PCT),
    (24, "Next-most-constrained stage", "=INDEX(B5:B11,MATCH(SMALL(K5:K11,2),K5:K11,0))", None),
    (25, "Next stage: utilisation at target", "=INDEX(G5:G11,MATCH(SMALL(K5:K11,2),K5:K11,0))", PCT),
]
for row, lab, fml, fmt in summ:
    put(ba, f"A{row}", None)
    ba.merge_cells(f"A{row}:C{row}")
    put(ba, f"A{row}", lab, BLACK, border=BOX)
    put(ba, f"D{row}", fml, BOLD, fmt, border=BOX, align=RIGHT)
    ba.merge_cells(f"D{row}:E{row}")
put(ba, "A27", "TOC rule: an hour lost at the constraint is an hour lost for the whole line; an hour saved at a non-constraint is a mirage.", SUB)

# ================================================================ FIXES
fx = wb.create_sheet("Fixes")
fx.sheet_view.showGridLines = False
put(fx, "A1", "Steps 2-5: deterministic scenario analysis (theoretical output, no variability)", TITLE)
put(fx, "A2", "Set 1/0 flags to switch fixes on. Levers, capacities, constraint and output recalculate. The constraint moves after each fix (step 5).", SUB)
heads = ["ID", "Scenario"] + S.FLAGS + ["Anodising lines", "Rack batch (min)", "Anodising availability", "Time multiplier", "Anodising yield",
                                           "Anodising capacity (units/wk)", "Line yield"] + \
    [None] * 7 + ["Max starts (starts/wk)", "Constraint after fix", "Max good output (units/wk)", "Sellable (capped at demand)", "Gain vs base (units/wk)",
                  "Starts needed for demand", "Headroom over demand", "One-time cost (USD)", "Annual cost (USD)"]
hdr(fx, 4, heads)
fx.row_dimensions[4].height = 58
for j in range(7):
    put(fx, f"{L(16 + j)}4", f"=BaseAnalysis!B{5 + j}", WHITE_B, fill=HDR, align=CENTER, border=BOX)
widths(fx, {"A": 6, "B": 38, **{L(c): 9 for c in range(3, 9)}, "I": 10, "J": 11, "K": 12, "L": 10, "M": 11, "N": 14, "O": 10,
            **{L(c): 12 for c in range(16, 23)}, "W": 13, "X": 24, "Y": 14, "Z": 14, "AA": 13, "AB": 13, "AC": 12, "AD": 13, "AE": 13})
SCN = S.SCENARIOS
F0 = 5
F1 = F0 + len(SCN) - 1
fcol = {name: L(3 + i) for i, name in enumerate(S.FLAGS)}  # DBR C, YIELD D, SMED E, SEVEN F, LINE3 G, TPM H
for i, (sid, name, flags) in enumerate(SCN):
    r_ = F0 + i
    put(fx, f"A{r_}", sid, BOLD, align=CENTER, border=BOX)
    put(fx, f"B{r_}", name, BLACK, border=BOX, align=WRAP)
    for fl in S.FLAGS:
        put(fx, f"{fcol[fl]}{r_}", 1 if fl in flags else 0, BLUE, '0', KEY, CENTER, BOX)
    put(fx, f"I{r_}", f"=IF(G{r_}=1,Inputs!$C$33,Inputs!$D$23)", BLACK, INT, align=CENTER, border=BOX)
    put(fx, f"J{r_}", f"=IF(E{r_}=1,Inputs!$C$31,Inputs!$F$23)", BLACK, N1, align=CENTER, border=BOX)
    put(fx, f"K{r_}", f"=IF(H{r_}=1,Inputs!$C$34,Inputs!$G$23)", BLACK, PCT, align=CENTER, border=BOX)
    put(fx, f"L{r_}", f"=IF(F{r_}=1,Inputs!$C$32,1)", BLACK, N3, align=CENTER, border=BOX)
    put(fx, f"M{r_}", f"=IF(D{r_}=1,Inputs!$C$30,Inputs!$I$23)", BLACK, PCT, align=CENTER, border=BOX)
    put(fx, f"N{r_}", f"=I{r_}*Inputs!$B$6*L{r_}*K{r_}*Inputs!$H$23/(J{r_}/Inputs!$B$9)", BLACK, INT, border=BOX)
    put(fx, f"O{r_}", f"=Inputs!$I$19*Inputs!$I$20*Inputs!$I$21*Inputs!$I$22*M{r_}*Inputs!$I$24*Inputs!$I$25", BLACK, PCT, align=CENTER, border=BOX)
    for j in range(4):
        put(fx, f"{L(16 + j)}{r_}", f"=BaseAnalysis!$K${5 + j}", GREEN, INT, border=BOX)
    put(fx, f"T{r_}", f"=N{r_}/BaseAnalysis!$D$9", BLACK, INT, border=BOX)
    put(fx, f"U{r_}", f"=BaseAnalysis!$F$10/(BaseAnalysis!$D$9*M{r_})", BLACK, INT, border=BOX)
    put(fx, f"V{r_}", f"=BaseAnalysis!$F$11/(BaseAnalysis!$D$9*M{r_}*Inputs!$I$24)", BLACK, INT, border=BOX)
    put(fx, f"W{r_}", f"=MIN(P{r_}:V{r_})", BOLD, INT, border=BOX)
    put(fx, f"X{r_}", f"=INDEX($P$4:$V$4,MATCH(W{r_},P{r_}:V{r_},0))", BOLD, border=BOX, align=WRAP)
    put(fx, f"Y{r_}", f"=W{r_}*O{r_}", BLACK, INT, border=BOX)
    put(fx, f"Z{r_}", f"=MIN(Y{r_},Inputs!$B$7)", BOLD, INT, border=BOX)
    put(fx, f"AA{r_}", f"=Z{r_}-$Z${F0}", BLACK, INT, border=BOX)
    put(fx, f"AB{r_}", f"=Inputs!$B$7/O{r_}", BLACK, INT, border=BOX)
    put(fx, f"AC{r_}", f"=Y{r_}/Inputs!$B$7-1", BLACK, PCT, border=BOX)
    put(fx, f"AD{r_}", "=" + "+".join(f"{fcol[fl]}{r_}*Inputs!$E${29 + k}" for k, fl in enumerate(S.FLAGS)), BLACK, USD0, border=BOX)
    put(fx, f"AE{r_}", "=" + "+".join(f"{fcol[fl]}{r_}*Inputs!$F${29 + k}" for k, fl in enumerate(S.FLAGS)), BLACK, USD0, border=BOX)
    fx.row_dimensions[r_].height = 30
fx.conditional_formatting.add(f"AC{F0}:AC{F1}", FormulaRule(formula=[f"AC{F0}<0"], fill=red))
fx.conditional_formatting.add(f"AC{F0}:AC{F1}", FormulaRule(formula=[f"AC{F0}>=0"], fill=grn))
fx.freeze_panes = "C5"

# ================================================================ SIM RESULTS
sr = wb.create_sheet("SimResults")
sr.sheet_view.showGridLines = False
widths(sr, {"A": 6, "B": 38, "C": 10, "D": 13, "E": 10, "F": 10, "G": 12, "H": 12, **{L(c): 11 for c in range(9, 17)}})
put(sr, "A1", "SimPy simulation results (pasted from simulate.py / run_experiments.py)", TITLE)
put(sr, "A2", f"SIMULATED DATA. {len(RES['seeds'])} replications per scenario (seeds {RES['seeds'][0]}-{RES['seeds'][-1]}), {RES['params']['weeks']} weeks with {RES['params']['warm_weeks']} warm-up. Rope size {RES['rope_chosen']} units. Blue = pasted outputs.", SUB)
hdr(sr, 4, ["ID", "Scenario", "Release", "Good units/wk (mean)", "Std dev", "95% CI +/-", "Avg WIP (units)", "Flow time (hours)", "Anodising busy %"] +
    [s["name"] for s in S.BASE_STAGES[:4]] + [s["name"] for s in S.BASE_STAGES[5:]])
put(sr, "Q4", "Rope used (units)", WHITE_B, fill=HDR, align=CENTER, border=BOX)
sr.column_dimensions["Q"].width = 12
sr.row_dimensions[4].height = 58
for i, sc in enumerate(RES["scenarios"]):
    r_ = 5 + i
    sm = sc["summary"]
    put(sr, f"A{r_}", sc["id"], BOLD, align=CENTER, border=BOX)
    put(sr, f"B{r_}", sc["name"], BLACK, border=BOX)
    put(sr, f"C{r_}", sc["policy"], BLACK, align=CENTER, border=BOX)
    put(sr, f"D{r_}", round(sm["good_wk"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"E{r_}", round(sm["good_wk"]["sd"], 1), BLUE, INT, border=BOX)
    put(sr, f"F{r_}", round(sm["good_wk"]["ci95"], 1), BLUE, INT, border=BOX)
    put(sr, f"G{r_}", round(sm["wip"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"H{r_}", round(sm["flow_hr"]["mean"], 2), BLUE, N1, border=BOX)
    put(sr, f"I{r_}", round(sm["util"][S.ANOD], 3), BLUE, PCT, border=BOX)
    put(sr, f"Q{r_}", sc.get("rope"), BLUE, INT, border=BOX)
    others = [u for k, u in enumerate(sm["util"]) if k != S.ANOD]
    for k, u in enumerate(others):
        put(sr, f"{L(10 + k)}{r_}", round(u, 3), BLUE, PCT, border=BOX)
SR1 = 5 + len(RES["scenarios"]) - 1
rs = SR1 + 3
sec(sr, rs - 1, "Rope (buffer) size sweep: drum-buffer-rope release with base-case capacity", 9)
hdr(sr, rs, ["", "Rope (units ahead of constraint)", "", "Good units/wk (mean)", "Std dev", "95% CI +/-", "Avg WIP (units)", "Flow time (hours)", "Anodising busy %"])
sr.row_dimensions[rs].height = 44
for i, sw in enumerate(RES["sweep"]):
    r_ = rs + 1 + i
    put(sr, f"B{r_}", sw["rope"], BLUE, INT, border=BOX)
    put(sr, f"D{r_}", round(sw["good_wk"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"E{r_}", round(sw["good_wk"]["sd"], 1), BLUE, INT, border=BOX)
    put(sr, f"F{r_}", round(sw["good_wk"]["ci95"], 1), BLUE, INT, border=BOX)
    put(sr, f"G{r_}", round(sw["wip"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"H{r_}", round(sw["flow_hr"]["mean"], 2), BLUE, N1, border=BOX)
    put(sr, f"I{r_}", round(sw["anod_util"], 3), BLUE, PCT, border=BOX)
SW0, SW1 = rs + 1, rs + len(RES["sweep"])
rs2 = SW1 + 5
sec(sr, rs2 - 1, "Rope re-sizing after the constraint is sped up (DBR + SMED scenario)", 9)
hdr(sr, rs2, ["", "Rope (units ahead of constraint)", "", "Good units/wk (mean)", "Std dev", "95% CI +/-", "Avg WIP (units)", "Flow time (hours)", "Anodising busy %"])
sr.row_dimensions[rs2].height = 44
for i, sw in enumerate(RES["sweep2"]):
    r_ = rs2 + 1 + i
    put(sr, f"B{r_}", sw["rope"], BLUE, INT, border=BOX)
    put(sr, f"D{r_}", round(sw["good_wk"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"E{r_}", round(sw["good_wk"]["sd"], 1), BLUE, INT, border=BOX)
    put(sr, f"F{r_}", round(sw["good_wk"]["ci95"], 1), BLUE, INT, border=BOX)
    put(sr, f"G{r_}", round(sw["wip"]["mean"], 1), BLUE, INT, border=BOX)
    put(sr, f"H{r_}", round(sw["flow_hr"]["mean"], 2), BLUE, N1, border=BOX)
    put(sr, f"I{r_}", round(sw["anod_util"], 3), BLUE, PCT, border=BOX)
put(sr, f"A{rs2 + len(RES['sweep2']) + 2}", "Lesson: a rope sized for the slow constraint starves it once capacity improves. Re-size the buffer after every exploit or elevate step (step 5).", SUB)
put(sr, f"A{SW1 + 2}", "Busy % excludes repair time. Under push release the anodising queue grows without bound when the constraint is overloaded, so WIP and flow time are time-averages over the measurement window and keep rising with run length.", SUB)

# ================================================================ ECONOMICS
ec = wb.create_sheet("Economics")
ec.sheet_view.showGridLines = False
widths(ec, {"A": 6, "B": 38, **{L(c): 14 for c in range(3, 19)}})
put(ec, "A1", "Gain versus cost (module level, USD)", TITLE)
put(ec, "A2", "Gain uses simulated output, capped at demand. Net benefit = annual contribution gain - annualised cost. Full-line column scales by the module count on Inputs.", SUB)
hdr(ec, 4, ["ID", "Scenario", "Sim good units/wk", "Sellable (capped)", "Gain vs base (units/wk)", "Gain (units/yr)", "Contribution gain (USD/yr)", "One-time cost (USD)",
            "Annual cost (USD)", "Annualised cost (USD/yr)", "Net annual benefit (USD)", "Payback (months)", "WIP vs base", "Flow time (h)", "Theory sellable (units/wk)",
            "Demand (units/wk)", "Rank by net benefit", "Full-line net benefit (USD/yr)"])
ec.row_dimensions[4].height = 58
for i in range(len(SCN)):
    r_ = F0 + i
    put(ec, f"A{r_}", f"=Fixes!A{r_}", GREEN, align=CENTER, border=BOX)
    put(ec, f"B{r_}", f"=Fixes!B{r_}", GREEN, border=BOX, align=WRAP)
    put(ec, f"C{r_}", f"=INDEX(SimResults!$D$5:$D${SR1},MATCH($A{r_},SimResults!$A$5:$A${SR1},0))", GREEN, INT, border=BOX)
    put(ec, f"D{r_}", f"=MIN(C{r_},Inputs!$B$7)", BLACK, INT, border=BOX)
    put(ec, f"E{r_}", f"=D{r_}-$D${F0}", BLACK, INT, border=BOX)
    put(ec, f"F{r_}", f"=E{r_}*Inputs!$B$8", BLACK, INT, border=BOX)
    put(ec, f"G{r_}", f"=F{r_}*Inputs!$B$13", BLACK, USD0, border=BOX)
    put(ec, f"H{r_}", f"=Fixes!AD{r_}", GREEN, USD0, border=BOX)
    put(ec, f"I{r_}", f"=Fixes!AE{r_}", GREEN, USD0, border=BOX)
    put(ec, f"J{r_}", f"=H{r_}/Inputs!$B$14+I{r_}", BLACK, USD0, border=BOX)
    put(ec, f"K{r_}", f"=G{r_}-J{r_}", BOLD, USD0, border=BOX)
    put(ec, f"L{r_}", f'=IF(H{r_}=0,0,IF(G{r_}-I{r_}>0,H{r_}/((G{r_}-I{r_})/12),"n/a"))', BLACK, N1, align=CENTER, border=BOX)
    put(ec, f"M{r_}", f"=INDEX(SimResults!$G$5:$G${SR1},MATCH($A{r_},SimResults!$A$5:$A${SR1},0))/SimResults!$G$5-1", BLACK, PCT, border=BOX)
    put(ec, f"N{r_}", f"=INDEX(SimResults!$H$5:$H${SR1},MATCH($A{r_},SimResults!$A$5:$A${SR1},0))", GREEN, N1, border=BOX)
    put(ec, f"O{r_}", f"=Fixes!Z{r_}", GREEN, INT, border=BOX)
    put(ec, f"P{r_}", "=Inputs!$B$7", GREEN, INT, border=BOX)
    put(ec, f"Q{r_}", f"=RANK(K{r_},$K${F0}:$K${F1},0)", BLACK, INT, align=CENTER, border=BOX)
    put(ec, f"R{r_}", f"=K{r_}*Inputs!$B$15", BLACK, USD0, border=BOX)
    ec.row_dimensions[r_].height = 30
ec.conditional_formatting.add(f"K{F0}:K{F1}", FormulaRule(formula=[f"K{F0}<0"], fill=red))
ec.conditional_formatting.add(f"K{F0}:K{F1}", FormulaRule(formula=[f"K{F0}>0"], fill=grn))
put(ec, f"A{F1 + 2}", "Reading the table: fixes at the constraint pay back fast; adding capacity (third line, 7-day working) costs more than the demand-capped gain once the cheap fixes exist. Recovered WIP and lead-time gains from DBR are real but are not converted to dollars here.", SUB)
ec.freeze_panes = "C5"

# ================================================================ DASHBOARD
db = wb.create_sheet("Dashboard")
db.sheet_view.showGridLines = False
widths(db, {"A": 58, "B": 40, "C": 3})
put(db, "A1", "Dashboard: finding and fixing the constraint on an enclosure production line", TITLE)
put(db, "A2", "SIMULATED / ILLUSTRATIVE DATA. Every figure is a formula (deterministic metrics) or a pasted SimPy result (SimResults).", SUB)
sec(db, 4, "Step 1: where is the constraint?", 2)
kp = [
    (5, "Constraint stage (base case)", "=BaseAnalysis!D19", None),
    (6, "Maximum good output per week (theory)", "=BaseAnalysis!D20", INT),
    (7, "Output as % of demand", "=BaseAnalysis!D22", PCT),
    (8, "Simulated base output (good units/wk)", "=SimResults!D5", INT),
    (9, "OEE of the constraint stage", "=BaseAnalysis!D23", PCT),
    (10, "Takt time per good unit (seconds)", "=BaseAnalysis!D16*60", N1),
    (11, "Next-most-constrained stage", "=BaseAnalysis!D24", None),
]
for row, lab, fml, fmt in kp:
    put(db, f"A{row}", lab, BLACK, border=BOX)
    put(db, f"B{row}", fml, BOLD, fmt, border=BOX, align=RIGHT)
sec(db, 13, "Steps 2-3: exploit and subordinate", 2)
kp2 = [
    (14, "DBR release: WIP reduction vs push (simulated)", "=-Economics!M6", PCT),
    (15, "Flow time, push release (hours)", "=SimResults!H5", N1),
    (16, "Flow time, DBR release (hours)", "=SimResults!H6", N1),
    (17, "SMED alone: simulated output (good units/wk)", "=Economics!C8", INT),
    (18, "SMED alone: net annual benefit (USD, module)", "=Economics!K8", USD0),
]
for row, lab, fml, fmt in kp2:
    put(db, f"A{row}", lab, BLACK, border=BOX)
    put(db, f"B{row}", fml, BOLD, fmt, border=BOX, align=RIGHT)
sec(db, 20, "Steps 4-5: elevate only if needed, then repeat", 2)
kp3 = [
    (21, "Best scenario by net annual benefit", f"=INDEX(Economics!B{F0}:B{F1},MATCH(1,Economics!Q{F0}:Q{F1},0))", None),
    (22, "Its net annual benefit (USD, module)", f"=INDEX(Economics!K{F0}:K{F1},MATCH(1,Economics!Q{F0}:Q{F1},0))", USD0),
    (23, "Low-cost package (DBR + SMED): net annual benefit (USD)", "=Economics!K12", USD0),
    (24, "Capex route (DBR + third line): net annual benefit (USD)", "=Economics!K15", USD0),
    (25, "Constraint after low-cost package", "=Fixes!X12", None),
    (26, "Low-cost package: headroom over demand (theory)", "=Fixes!AC12", PCT),
    (27, "Low-cost package: full-line net benefit (USD/yr)", "=Economics!R12", USD0),
]
for row, lab, fml, fmt in kp3:
    put(db, f"A{row}", lab, BLACK, border=BOX)
    put(db, f"B{row}", fml, BOLD, fmt, border=BOX, align=RIGHT)

ch1 = BarChart()
ch1.type = "bar"
ch1.style = 10
ch1.title = "Utilisation at target demand (above 100% = constraint)"
ch1.add_data(Reference(ba, min_col=7, max_col=7, min_row=4, max_row=11), titles_from_data=True)
ch1.set_categories(Reference(ba, min_col=2, min_row=5, max_row=11))
ch1.legend = None
ch1.gapWidth = 50
ch1.y_axis.number_format = "0%"
ch1.y_axis.scaling.min = 0
ch1.x_axis.scaling.orientation = "maxMin"
ch1.y_axis.delete = False
ch1.x_axis.delete = False
ch1.height, ch1.width = 9, 20
db.add_chart(ch1, "D4")

ch2 = BarChart()
ch2.type = "col"
ch2.style = 10
ch2.title = "Good output per week: simulated vs theory vs demand"
ch2.add_data(Reference(ec, min_col=4, max_col=4, min_row=4, max_row=F1), titles_from_data=True)
ch2.add_data(Reference(ec, min_col=15, max_col=16, min_row=4, max_row=F1), titles_from_data=True)
ch2.set_categories(Reference(ec, min_col=1, min_row=F0, max_row=F1))
ch2.y_axis.scaling.min = 8000
ch2.y_axis.scaling.max = 10500
ch2.gapWidth = 60
ch2.y_axis.delete = False
ch2.x_axis.delete = False
ch2.height, ch2.width = 9, 20
db.add_chart(ch2, "D23")

ch3 = BarChart()
ch3.type = "col"
ch3.style = 10
ch3.title = "Net annual benefit by scenario (USD per module)"
ch3.add_data(Reference(ec, min_col=11, max_col=11, min_row=4, max_row=F1), titles_from_data=True)
ch3.set_categories(Reference(ec, min_col=1, min_row=F0, max_row=F1))
ch3.legend = None
ch3.gapWidth = 50
ch3.y_axis.number_format = '$#,##0'
ch3.x_axis.tickLblPos = "low"
ch3.y_axis.delete = False
ch3.x_axis.delete = False
ch3.height, ch3.width = 9, 20
db.add_chart(ch3, "D42")

ch4 = LineChart()
ch4.style = 12
ch4.title = "DBR rope size: good units per week"
ch4.add_data(Reference(sr, min_col=4, max_col=4, min_row=rs, max_row=SW1), titles_from_data=True)
ch4.set_categories(Reference(sr, min_col=2, min_row=SW0, max_row=SW1))
ch4.y_axis.scaling.min = 7800
ch4.y_axis.scaling.max = 9200
ch4.legend = None
ch4.x_axis.title = "Rope (units ahead of constraint)"
ch4.y_axis.delete = False
ch4.x_axis.delete = False
ch4.height, ch4.width = 9, 20
db.add_chart(ch4, "D61")

ch5 = BarChart()
ch5.type = "col"
ch5.style = 10
ch5.title = "Average WIP by scenario (units, simulated)"
ch5.add_data(Reference(sr, min_col=7, max_col=7, min_row=4, max_row=SR1), titles_from_data=True)
ch5.set_categories(Reference(sr, min_col=1, min_row=5, max_row=SR1))
ch5.legend = None
ch5.gapWidth = 50
ch5.y_axis.delete = False
ch5.x_axis.delete = False
ch5.height, ch5.width = 9, 20
db.add_chart(ch5, "D80")

# ================================================================ SOURCES
so = wb.create_sheet("Sources")
so.sheet_view.showGridLines = False
widths(so, {"A": 6, "B": 62, "C": 70, "D": 80})
put(so, "A1", "Sources and basis", TITLE)
put(so, "A2", "The line is simulated. Benchmarks inform plausible ranges only; nothing here describes a real plant.", SUB)
hdr(so, 4, ["Ref", "Source", "Used for", "URL / reference"])
srcs = [
    ("W1", "Fabrico: OEE benchmark guide (world-class 85% = 90% availability x 95% performance x 99.9% quality; typical about 60%)", "Plausible availability and performance ranges for stages", "https://www.fabrico.io/blog/oee-benchmark/"),
    ("W2", "TeepTrak: world-class OEE benchmark (metal fabrication and general machining: world-class 80-85%, average 55-65%)", "Sanity-check on machining-stage OEE", "https://teeptrak.com/en/world-class-oee-benchmark/"),
    ("W3", "Tractian: what is world-class OEE (Nakajima TPM benchmarks)", "OEE definition: availability x performance x quality", "https://tractian.com/en/blog/world-class-oee"),
    ("W4", "SimPy documentation: discrete-event simulation in Python", "Simulation framework", "https://simpy.readthedocs.io/"),
    ("B1", "Goldratt and Cox, The Goal (1984); Goldratt, Theory of Constraints and the Five Focusing Steps; drum-buffer-rope scheduling", "Method: identify, exploit, subordinate, elevate, repeat; DBR release", "Book reference (no URL)"),
    ("P1", "Project 1: TCO and scenario model, aluminium enclosure (India base case)", "Ex-works price, net material cost, and the 86% first-pass yield used as the base line yield", "Apple_Sourcing_TCO_Model_Aluminium_Enclosure.xlsx"),
    ("P2", "Project 2: localization gap and supplier risk dashboard", "Context: capacity ramp and yield are top-three risks for Indian enclosure suppliers", "Enclosure_Localization_Supplier_Risk_Dashboard.xlsx"),
    ("A1", "ASSUMPTION: all stage cycle times, machine counts, MTTR, fix parameters and fix costs", "Illustrative values; replace with time studies and supplier quotes", "n/a"),
]
for i, s_ in enumerate(srcs):
    for j, v in enumerate(s_):
        put(so, f"{L(1 + j)}{5 + i}", v, BOLD if (j == 0 and s_[0] == "A1") else BLACK, align=WRAP, border=BOX)
    so.row_dimensions[5 + i].height = 34

for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

order = ["README", "Dashboard", "Inputs", "BaseAnalysis", "Fixes", "SimResults", "Economics", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("saved", OUT, "P1 price", P1_PRICE, "material", P1_MATERIAL)
