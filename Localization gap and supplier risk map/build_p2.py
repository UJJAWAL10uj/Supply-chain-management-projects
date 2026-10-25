"""
Project 2: Localization gap and supplier risk map for mobile aluminium enclosures in India.
All scores, ranks, ratings, KPIs and heat-map counts are Excel formulas. Run: python build_p2.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.properties import PageSetupProperties

OUT = "/mnt/user-data/outputs/Enclosure_Localization_Supplier_Risk_Dashboard.xlsx"
FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
WHITE_B = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=10, italic=True, color="595959")
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
KEY = PatternFill("solid", fgColor="FFFF00")
TOT = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="top")
NUM1 = '0.0;(0.0);"-"'
NUM2 = '0.00;(0.00);"-"'
INT = '0;(0);"-"'
PCT = '0%;(0%);"-"'

wb = Workbook()


def put(ws, ref, v, font=BLACK, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def hdr(ws, row, labels, c0=1):
    for i, t in enumerate(labels):
        put(ws, f"{L(c0 + i)}{row}", t, WHITE_B, fill=HDR, align=CENTER, border=BOX)


def sec(ws, row, text, n):
    for i in range(1, n + 1):
        ws[f"{L(i)}{row}"].fill = SEC
    put(ws, f"A{row}", text, BOLD, fill=SEC)


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


# =========================================================== README
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 26, "C": 112})
put(ws, "B2", "Localization Gap and Supplier Risk Map: Mobile Aluminium Enclosures in India", TITLE)
put(ws, "B3", "Independent academic project built from public reports. Scores are analyst judgments, not Apple data.", SUB)
rows = [
    ("Question", "Which parts of the mobile aluminium-enclosure value chain can India make today versus import, which Indian suppliers carry the most risk, and what should a sourcing team do about the top three risks?"),
    ("Category", "Mobile / consumer-electronics aluminium enclosures (metal housings and chassis). Chosen because enclosures are the most heavily incentivised component category under India's ECMS and the one with the largest visible import gap (extruded aluminium)."),
    ("Sheets", "Dashboard (headline metrics, rankings, charts) > ValueChain (localization gap by stage) > Suppliers (10 suppliers scored on four criteria) > Risks (register, ranking, heat map) > Mitigation (actions for the top three risks) > Sources."),
    ("Scoring logic", "Supplier risk score = weighted average of four 1-5 scores (5 = highest risk): single-source exposure, capacity and ramp risk, geopolitical / China-linked exposure, quality and process readiness. Weights are editable. Value-chain gap score = import dependence (1-5) x (6 - India readiness (1-5)), range 1-25. Risk score = likelihood x impact (1-5 each)."),
    ("How to use", "Change only BLUE cells (scores, weights, thresholds, likelihood, impact). Ranks, ratings, KPIs, charts and the heat map recalculate. Yellow = key lever."),
    ("Colour code", "Blue = input. Black = formula. Green = link to another sheet. Yellow fill = key assumption."),
    ("Evidence", "Facts about suppliers, ECMS approvals and incidents come from the public reports on the Sources sheet and are cited as [n]. The 1-5 scores are my interpretation of those facts. Rows marked 'Analyst estimate' on ValueChain have no direct public source and should be verified."),
    ("Limitations", "Press reports can be dated or based on unnamed sources. Plant status changes quickly (check before relying on any single item). Supplier scores compare suppliers with each other; they are not statements about any company's absolute quality."),
]
r = 5
for k, v in rows:
    put(ws, f"B{r}", k, BOLD, align=WRAP)
    put(ws, f"C{r}", v, BLACK, align=WRAP)
    ws.row_dimensions[r].height = 62 if len(v) > 230 else 46
    r += 1

# =========================================================== SUPPLIERS
sp = wb.create_sheet("Suppliers")
sp.sheet_view.showGridLines = False
widths(sp, {"A": 6, "B": 30, "C": 26, "D": 11, "E": 22, "F": 24, "G": 70, "H": 11, "I": 11, "J": 13, "K": 11, "L": 12, "M": 10, "N": 7, "O": 80})
put(sp, "A1", "Supplier risk scoring: mobile aluminium enclosure chain in India", TITLE)
put(sp, "A2", "Scores 1-5 where 5 = highest risk. Weights and thresholds are editable. Scores are analyst judgments based on the cited reports.", SUB)
hdr(sp, 4, ["Criterion", "Single-source exposure", "Capacity and ramp risk", "Geopolitical / China-linked exposure", "Quality and process readiness", "Total"], 1)
sp.merge_cells("A4:A4")
put(sp, "A5", "Weight", BOLD, border=BOX)
for j, w in enumerate([0.25, 0.25, 0.25, 0.25]):
    put(sp, f"{L(2 + j)}5", w, BLUE, PCT, KEY, border=BOX)
put(sp, "F5", "=SUM(B5:E5)", BLACK, PCT, border=BOX)
put(sp, "G5", '=IF(ROUND(F5,6)=1,"Weights sum to 100%","Weights do not sum to 100%: scores are normalised")', SUB)
put(sp, "A6", "Medium from", BOLD)
put(sp, "B6", 2.5, BLUE, NUM1, KEY, border=BOX)
put(sp, "C6", "score (below = Low)", SUB)
put(sp, "A7", "High from", BOLD)
put(sp, "B7", 3.5, BLUE, NUM1, KEY, border=BOX)
put(sp, "C7", "score", SUB)
hdr(sp, 9, ["ID", "Supplier", "Role in enclosure chain", "Tier", "Location", "Owner / HQ", "Apple link and public status (sources in [ ])",
            "Single-source", "Capacity / ramp", "Geopolitical", "Quality readiness", "Weighted risk score", "Rating", "Rank", "Why these scores (analyst judgment on the evidence)"])
sp.row_dimensions[9].height = 44

sup = [
    ("S01", "Tata Electronics (Hosur)", "iPhone metal enclosures (CNC, finishing)", "Tier-1", "Hosur, Tamil Nadu", "Tata Group (India)",
     "Makes iPhone enclosures at Hosur; Foxconn's later entry ended Tata's position as the only local maker [19]. Supplies only a limited share of iPhone enclosures domestically, set to rise [9]. ECMS enclosure approval Jan 2026; Hosur enclosure plant moving to production Aug 2026 [4][6]. Plant restored after a Sept 2024 fire; output reportedly being doubled [2].",
     4, 3, 3, 2,
     "Single-source 4: long the only local maker and still a limited share of volume [9][19]. Capacity 3: the 2024 fire shows plant-level fragility [2]. Geopolitical 3: Indian-owned, but extrusion is 100% imported and tooling largely Chinese [4][10]. Quality 2: producing iPhone enclosures at volume for several years [18]."),
    ("S02", "Yuzhan Technology (India)", "iPhone enclosures and display modules", "Tier-1", "Kanchipuram, Tamil Nadu", "Foxconn (Taiwan)",
     "Foxconn's India arm; made enclosures and display modules for older iPhones, not yet iPhone 17, as of mid-2025 [11]. ECMS enclosure approval Jan 2026; Kanchipuram enclosure facility progressing to production (Aug 2026) [4][6]. Foxconn recalled about 300 Chinese engineers from India in 2025 [11].",
     3, 4, 5, 4,
     "Single-source 3: second local iPhone enclosure source. Capacity 4: production began only months before mid-2025 and a new plant is still being built [6][11]. Geopolitical 5: engineer recall and held-up equipment shipments hit Foxconn directly [10][11]. Quality 4: not on the latest iPhone line as of mid-2025 [11]."),
    ("S03", "Motherson Electronic Components", "Enclosures / mechanical components", "Tier-1", "Kanchipuram, Tamil Nadu", "Motherson Group (India)",
     "ECMS enclosure approval Jan 2026; Kanchipuram plant expected to create 5,741 jobs [4][6][20]. Reported to produce mechanical components such as enclosures for Apple products, shipped to China and Vietnam assembly [1].",
     2, 4, 2, 3,
     "Single-source 2: an additional rather than sole source. Capacity 4: large plant still being built [6][20]. Geopolitical 2: Indian group. Quality 3: exporting enclosure parts for several Apple products but metal-enclosure volume not yet proven [1]."),
    ("S04", "Aequs (Aequs Consumer Products)", "MacBook, Apple Watch and other non-iPhone enclosures", "Tier-1", "Hubballi, Karnataka", "Aequs (India)",
     "Trial production of MacBook enclosures and Apple Watch mechanical parts; formally added to Apple's supply chain [3]. ECMS approval for laptop and smartwatch enclosures [5]. Electronics SEZ cleared June 2025 [3].",
     2, 3, 2, 3,
     "Single-source 2: niche, non-iPhone parts. Capacity 3: trial production, SEZ still being developed [3]. Geopolitical 2: Indian company. Quality 3: qualification under way, volume history limited [3]."),
    ("S05", "Jabil India", "AirPods and plastic enclosures", "Tier-1", "Pune, Maharashtra; Trichy, Tamil Nadu (planned)", "Jabil Inc. (US)",
     "Produces AirPods components at Pune; may take on plastic AirPods enclosures [1][21]. Announced a plant near Trichy of about ₹2,000 crore ($238M), around 5,000 jobs [14].",
     2, 3, 2, 2,
     "Single-source 2: global multi-site firm. Capacity 3: Trichy plant is planned, not yet running [14]. Geopolitical 2: US parent. Quality 2: established global Apple supplier [1]."),
    ("S06", "TD Connex", "Micro-precision CNC, injection, stamping, LSR", "Tier-2", "Oragadam, Tamil Nadu", "Singapore-based",
     "Operates a unit at Oragadam; preparing a ₹200 crore expansion and about 20 acres for smartphone-housing functions (CNC, injection, stamping, LSR) [3][15].",
     3, 3, 3, 3,
     "Single-source 3: small niche supplier with few alternates. Capacity 3: expansion planned, land not yet acquired [3]. Geopolitical 3: foreign-owned, ownership structure not verified [15]. Quality 3: capabilities listed but Apple volume history not verified."),
    ("S07", "Hindalco Industries (extrusion)", "Precision aluminium extrusion for enclosures", "Tier-2 (material)", "Kuppam, Andhra Pradesh; Silvassa", "Aditya Birla Group (India)",
     "Phone-enclosure extrusion is 100% imported today [4][8]. ₹586 crore Kuppam plant for iPhone chassis, production targeted March 2027 [7]. ECMS approval Jan 2026 [4]. Extrusion capacity to rise from 60,000 to 109,000 tonnes a year with Kuppam and Silvassa [8].",
     5, 4, 2, 4,
     "Single-source 5: the only planned domestic source of a fully imported input [4][8]. Capacity 4: plant not producing until March 2027 [7]. Geopolitical 2: Indian metal. Quality 4: phone-grade precision extrusion is new and still to be qualified."),
    ("S08", "Titan Engineering & Automation", "Capital equipment and automation", "Tier-2 (equipment)", "Hosur, Tamil Nadu", "Titan Company (Tata Group)",
     "One of the Indian firms supplying capital equipment to Apple's vendors; around 35 Indian firms make machinery for Apple's local supply chain [12].",
     3, 3, 2, 3,
     "Single-source 3: one of about 35 machinery makers [12]. Capacity 3: scaling with Apple vendors. Geopolitical 2: Indian company, reduces China tooling exposure. Quality 3: precision consumer-electronics tooling is a new domain."),
    ("S09", "Jyoti CNC Automation", "CNC machine tools", "Tier-2 (equipment)", "Rajkot, Gujarat", "Jyoti CNC (India)",
     "Named among Indian firms supplying capital equipment to Apple's vendors as China restricts equipment exports to them [12].",
     3, 3, 2, 4,
     "Single-source 3: one of about 35 makers [12]. Capacity 3. Geopolitical 2: Indian. Quality 4: local machines unproven at small-format consumer-electronics tolerances; Rest of World reports India lacked the technology for some specialised equipment [10]."),
    ("S10", "Bharat Forge", "Precision engineering / equipment", "Tier-2 (equipment)", "Pune, Maharashtra", "Bharat Forge (India)",
     "Named among Indian firms supplying capital equipment to Apple's vendors [12].",
     3, 2, 2, 4,
     "Single-source 3: one of about 35 makers [12]. Capacity 2: large industrial group. Geopolitical 2: Indian. Quality 4: moving from heavy engineering to consumer-electronics precision is a new domain."),
]
S0, S1 = 10, 10 + len(sup) - 1
for i, row in enumerate(sup):
    r_ = S0 + i
    for j in range(7):
        put(sp, f"{L(1 + j)}{r_}", row[j], BLACK, align=WRAP, border=BOX)
    for j in range(4):
        put(sp, f"{L(8 + j)}{r_}", row[7 + j], BLUE, INT, KEY, CENTER, BOX)
    put(sp, f"L{r_}", f"=SUMPRODUCT(H{r_}:K{r_},$B$5:$E$5)/$F$5", BOLD, NUM2, align=CENTER, border=BOX)
    put(sp, f"M{r_}", f'=IF(L{r_}>=$B$7,"High",IF(L{r_}>=$B$6,"Medium","Low"))', BOLD, align=CENTER, border=BOX)
    put(sp, f"N{r_}", f"=RANK(L{r_},$L${S0}:$L${S1},0)+COUNTIF($L${S0}:L{r_},L{r_})-1", BLACK, INT, align=CENTER, border=BOX)
    put(sp, f"O{r_}", row[11], BLACK, align=WRAP, border=BOX)
    sp.row_dimensions[r_].height = 92
red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
amb = PatternFill("solid", bgColor="FFE699", fgColor="FFE699")
grn = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
sp.conditional_formatting.add(f"M{S0}:M{S1}", FormulaRule(formula=[f'M{S0}="High"'], fill=red))
sp.conditional_formatting.add(f"M{S0}:M{S1}", FormulaRule(formula=[f'M{S0}="Medium"'], fill=amb))
sp.conditional_formatting.add(f"M{S0}:M{S1}", FormulaRule(formula=[f'M{S0}="Low"'], fill=grn))
sp.freeze_panes = "C10"

# =========================================================== VALUE CHAIN
vc = wb.create_sheet("ValueChain")
vc.sheet_view.showGridLines = False
widths(vc, {"A": 6, "B": 34, "C": 40, "D": 70, "E": 12, "F": 12, "G": 10, "H": 8, "I": 13, "J": 17, "K": 38, "L": 12})
put(vc, "A1", "Localization gap map: what India makes today versus imports", TITLE)
put(vc, "A2", "Import dependence 1-5 (5 = fully imported). India readiness 1-5 (5 = mature, scaled capacity). Gap score = dependence x (6 - readiness).", SUB)
put(vc, "A3", "Critical from", BOLD)
put(vc, "B3", 15, BLUE, INT, KEY, border=BOX)
put(vc, "A4", "Elevated from", BOLD)
put(vc, "B4", 9, BLUE, INT, KEY, border=BOX)
hdr(vc, 6, ["ID", "Value-chain stage", "What it is", "India status today (public evidence in [ ])", "Import dependence", "India readiness", "Gap score", "Rank", "Priority", "Evidence strength", "Key Indian players / projects", "Source ref"])
vc.row_dimensions[6].height = 36
stages = [
    ("V1", "Primary aluminium and alloy", "Smelted primary metal and 6xxx-series alloy billet",
     "India has large integrated aluminium producers; Hindalco plans a $10bn five-year expansion including the Dahej smelter. Metal supply is not the constraint; precision extrusion is.",
     2, 4, "Mixed", "Hindalco, Vedanta, NALCO", "[8][22]"),
    ("V2", "Precision extrusion for enclosures", "Extruded aluminium profile later machined into the enclosure",
     "100% imported today per the government's own ECMS release. Hindalco's ₹586 crore Kuppam plant targets March 2027; extrusion capacity to rise 60,000 to 109,000 tonnes a year with Kuppam and Silvassa.",
     5, 2, "Sourced", "Hindalco (Kuppam, Silvassa)", "[4][7][8]"),
    ("V3", "CNC machining of enclosures", "Precision milling of the extruded or billet blank into the final enclosure",
     "Tata supplies a limited share of iPhone enclosures; Yuzhan, Motherson, Aequs and TD Connex are ramping. Indian suppliers already export mechanical parts and housings to China and Vietnam.",
     3, 3, "Sourced", "Tata Electronics, Yuzhan, Motherson, Aequs, TD Connex", "[3][6][9][13]"),
    ("V4", "CNC machine tools and precision equipment", "The machines and tooling that cut and finish enclosures",
     "Equipment shipments to Foxconn's India plants were held up in China in 2025; India lacked technology for some specialised equipment. About 35 Indian firms now make machinery for Apple's supply chain, nearly half engaged in the last two years.",
     4, 2, "Sourced", "Titan Engineering & Automation, Jyoti CNC, Bharat Forge, Wipro", "[10][11][12]"),
    ("V5", "Cutting tools, fixtures and consumables", "Cutters, jigs, fixtures and wear parts consumed in machining",
     "No direct public source found. Likely tied to the same China-linked equipment ecosystem as V4; treat as an estimate and verify with suppliers.",
     4, 3, "Analyst estimate", "Not identified in sources", "n/a"),
    ("V6", "Surface finishing (anodising and chemistry)", "Anodising, dyeing and coating that give the enclosure its finish",
     "No direct public source found. Process chemistry and dyes may rely on imports; treat as an estimate and verify with finishing suppliers.",
     3, 3, "Analyst estimate", "Not identified in sources", "n/a"),
    ("V7", "Plastic and antenna-band moulding, LSR sealing", "Injection-moulded antenna splits and liquid silicone seals inside the housing",
     "TD Connex lists plastic injection, liquid silicone rubber and stamping capability at Oragadam; Jabil and Motherson produce plastic and mechanical enclosure parts.",
     2, 3, "Mixed", "TD Connex, Jabil, Motherson", "[1][3][15]"),
    ("V8", "Process know-how and skilled engineers", "Line-setup, tool-maintenance and process-engineering skills",
     "Foxconn recalled about 300 Chinese engineers from India in 2025 and began flying in Taiwanese replacements; Beijing reportedly discouraged technology and equipment transfers to India.",
     4, 2, "Sourced", "Foxconn, Tata Electronics (training programmes)", "[11][23]"),
    ("V9", "Inspection, metrology and automation", "Vision inspection, gauging and line automation",
     "Titan Engineering & Automation and VVDN appear on the reported list of Apple suppliers in India; no capability detail found. Treat as an estimate.",
     3, 3, "Analyst estimate", "Titan Engineering & Automation, VVDN", "[12][24]"),
]
V0, V1 = 7, 7 + len(stages) - 1
for i, s in enumerate(stages):
    r_ = V0 + i
    put(vc, f"A{r_}", s[0], BLACK, align=WRAP, border=BOX)
    put(vc, f"B{r_}", s[1], BOLD, align=WRAP, border=BOX)
    put(vc, f"C{r_}", s[2], BLACK, align=WRAP, border=BOX)
    put(vc, f"D{r_}", s[3], BLACK, align=WRAP, border=BOX)
    put(vc, f"E{r_}", s[4], BLUE, INT, KEY, CENTER, BOX)
    put(vc, f"F{r_}", s[5], BLUE, INT, KEY, CENTER, BOX)
    put(vc, f"G{r_}", f"=E{r_}*(6-F{r_})", BOLD, INT, align=CENTER, border=BOX)
    put(vc, f"H{r_}", f"=RANK(G{r_},$G${V0}:$G${V1},0)+COUNTIF($G${V0}:G{r_},G{r_})-1", BLACK, INT, align=CENTER, border=BOX)
    put(vc, f"I{r_}", f'=IF(G{r_}>=$B$3,"Critical",IF(G{r_}>=$B$4,"Elevated","Manageable"))', BOLD, align=CENTER, border=BOX)
    put(vc, f"J{r_}", s[6], BLACK, align=CENTER, border=BOX)
    put(vc, f"K{r_}", s[7], BLACK, align=WRAP, border=BOX)
    put(vc, f"L{r_}", s[8], BLACK, align=CENTER, border=BOX)
    vc.row_dimensions[r_].height = 78
vc.conditional_formatting.add(f"I{V0}:I{V1}", FormulaRule(formula=[f'I{V0}="Critical"'], fill=red))
vc.conditional_formatting.add(f"I{V0}:I{V1}", FormulaRule(formula=[f'I{V0}="Elevated"'], fill=amb))
vc.conditional_formatting.add(f"I{V0}:I{V1}", FormulaRule(formula=[f'I{V0}="Manageable"'], fill=grn))
vc.freeze_panes = "C7"

# =========================================================== RISKS
rk = wb.create_sheet("Risks")
rk.sheet_view.showGridLines = False
widths(rk, {"A": 6, "B": 36, "C": 74, "D": 26, "E": 11, "F": 11, "G": 9, "H": 8, "I": 10, "J": 44, "K": 11, "L": 11, "M": 10, "N": 11})
put(rk, "A1", "Risk register: mobile aluminium enclosure supply chain in India", TITLE)
put(rk, "A2", "Likelihood and impact scored 1-5 (analyst judgment). Score = likelihood x impact. Top three are tagged; residual scores assume the Mitigation sheet actions are done.", SUB)
hdr(rk, 4, ["ID", "Risk", "Description and evidence (sources in [ ])", "Linked stage / supplier", "Likelihood", "Impact", "Score", "Rank", "Top 3?", "Early-warning indicator", "Residual likelihood", "Residual impact", "Residual score", "Reduction"])
rk.row_dimensions[4].height = 44
risks = [
    ("R1", "Imported aluminium extrusion dependence", "Extruded aluminium for phone enclosures is 100% imported until Hindalco's Kuppam plant starts (target March 2027). Any trade, tariff or logistics shock hits every enclosure maker at once [4][7][8].", "V2; S07 Hindalco", 4, 5, "Weeks of extrusion cover; Kuppam commissioning milestones; supplier lead times", 3, 3),
    ("R2", "Chinese equipment and engineer dependence", "Foxconn recalled about 300 Chinese engineers and equipment shipments were held up in China in 2025; India lacked technology for some specialised equipment [10][11][23].", "V4, V8; S02 Yuzhan, S08-S10", 4, 5, "Equipment shipment delays; engineer headcount; share of equipment value from non-China origin", 3, 4),
    ("R3", "Capacity ramp and execution delay", "Several new enclosure plants (Hindalco Kuppam, Motherson and Yuzhan at Kanchipuram, Tata Hosur) must ramp together; one slip can stall launch volumes [4][6][7].", "V3; S01, S02, S03, S07", 4, 4, "Schedule slip in weeks; line OEE; first-pass yield versus plan", 3, 3),
    ("R4", "Single-site concentration and plant incident", "Tata's Hosur plant is the main local iPhone enclosure source and had a fire in Sept 2024; one site outage removes most local supply [2][9].", "V3; S01 Tata Electronics", 3, 5, "Safety audit findings; days of finished-goods buffer", None, None),
    ("R5", "Yield and quality ramp at new Indian lines", "Early lines run below mature-line yield; the TCO model shows India needs about 95% first-pass yield to match China's landed cost at 50% incentive pass-through (Project 1).", "V3; S02, S03, S04", 4, 3, "Weekly first-pass yield; scrap rate; customer returns", None, None),
    ("R6", "Policy and incentive risk", "ECMS payouts, duty rates and US tariff policy can change; ECMS approvals already exceed the scheme's original target of ₹59,350 crore [6].", "All", 3, 3, "Notification changes; ECMS disbursement delays", None, None),
    ("R7", "Skilled labour hiring and attrition", "Tata alone planned to add up to 45,000 workers at Hosur; hiring and retention at that scale is an execution risk [18].", "V8; S01", 3, 3, "Attrition rate; training completion rate", None, None),
    ("R8", "Aluminium price volatility", "LME cash aluminium moved from $3,248.5 to $3,107 per tonne in one week (28 Sep to 5 Oct 2026) [25]; analysts see $3,200-3,500 average for 2026 [26].", "V1; all", 4, 2, "LME price versus contract index band", None, None),
]
R0, R1 = 5, 5 + len(risks) - 1
for i, s in enumerate(risks):
    r_ = R0 + i
    put(rk, f"A{r_}", s[0], BLACK, align=WRAP, border=BOX)
    put(rk, f"B{r_}", s[1], BOLD, align=WRAP, border=BOX)
    put(rk, f"C{r_}", s[2], BLACK, align=WRAP, border=BOX)
    put(rk, f"D{r_}", s[3], BLACK, align=WRAP, border=BOX)
    put(rk, f"E{r_}", s[4], BLUE, INT, KEY, CENTER, BOX)
    put(rk, f"F{r_}", s[5], BLUE, INT, KEY, CENTER, BOX)
    put(rk, f"G{r_}", f"=E{r_}*F{r_}", BOLD, INT, align=CENTER, border=BOX)
    put(rk, f"H{r_}", f"=RANK(G{r_},$G${R0}:$G${R1},0)+COUNTIF($G${R0}:G{r_},G{r_})-1", BLACK, INT, align=CENTER, border=BOX)
    put(rk, f"I{r_}", f'=IF(H{r_}<=3,"TOP 3","")', BOLD, align=CENTER, border=BOX)
    put(rk, f"J{r_}", s[6], BLACK, align=WRAP, border=BOX)
    if s[7] is not None:
        put(rk, f"K{r_}", s[7], BLUE, INT, KEY, CENTER, BOX)
        put(rk, f"L{r_}", s[8], BLUE, INT, KEY, CENTER, BOX)
    else:
        put(rk, f"K{r_}", None, BLACK, border=BOX)
        put(rk, f"L{r_}", None, BLACK, border=BOX)
    put(rk, f"M{r_}", f'=IF(AND(ISNUMBER(K{r_}),ISNUMBER(L{r_})),K{r_}*L{r_},"")', BOLD, INT, align=CENTER, border=BOX)
    put(rk, f"N{r_}", f'=IF(ISNUMBER(M{r_}),1-M{r_}/G{r_},"")', BLACK, PCT, align=CENTER, border=BOX)
    rk.row_dimensions[r_].height = 66
rk.conditional_formatting.add(f"I{R0}:I{R1}", FormulaRule(formula=[f'I{R0}="TOP 3"'], fill=red))

HR = R1 + 4  # heat-map header row
sec(rk, HR - 1, "Risk heat map: number of risks in each likelihood x impact cell", 14)
put(rk, f"D{HR}", "Impact \\ Likelihood", WHITE_B, fill=HDR, align=CENTER, border=BOX)
for j in range(5):
    put(rk, f"{L(5 + j)}{HR}", j + 1, WHITE_B, '0', HDR, CENTER, BOX)
for i, imp in enumerate([5, 4, 3, 2, 1]):
    rr = HR + 1 + i
    put(rk, f"D{rr}", imp, BOLD, '"Impact "0', align=CENTER, border=BOX)
    for j in range(5):
        col = L(5 + j)
        put(rk, f"{col}{rr}", f"=COUNTIFS($E${R0}:$E${R1},{col}${HR},$F${R0}:$F${R1},$D{rr})", BOLD, INT, align=CENTER, border=BOX)
HM = f"E{HR + 1}:I{HR + 5}"
rk.conditional_formatting.add(HM, FormulaRule(formula=[f"E${HR}*$D{HR + 1}>=15"], fill=red))
rk.conditional_formatting.add(HM, FormulaRule(formula=[f"E${HR}*$D{HR + 1}>=8"], fill=amb))
rk.conditional_formatting.add(HM, FormulaRule(formula=[f"E${HR}*$D{HR + 1}<8"], fill=grn))
rk.freeze_panes = "C5"

# =========================================================== MITIGATION
mt = wb.create_sheet("Mitigation")
mt.sheet_view.showGridLines = False
widths(mt, {"A": 7, "B": 34, "C": 5, "D": 76, "E": 26, "F": 14, "G": 46, "H": 30})
put(mt, "A1", "Mitigation plan for the top three risks", TITLE)
put(mt, "A2", "Actions are recommendations written from the evidence; owners are functions, not named people. Risk names pull from the Risks sheet. If you change the top three, update the action rows.", SUB)
hdr(mt, 4, ["Risk", "Risk name", "#", "Action", "Owner (function)", "Timing", "KPI / trigger", "Target"])
mit = [
    ("R1", 1, "Stage-gate qualification of Hindalco Kuppam extrusion: sample lots, dimensional and surface capability studies, then volume release tied to measured yield, ahead of the March 2027 start.", "Sourcing + Supplier quality", "0-9 months", "Qualified lots passed; domestic extrusion share of total", "Kuppam qualified before volume release"),
    ("R1", 2, "Hold safety stock and keep two approved import sources for extrusion until domestic yield is proven; set weeks of cover by lead time and lane risk.", "Supply chain planning", "0-3 months", "Weeks of extrusion cover", "At least the longest-lane lead time plus buffer"),
    ("R1", 3, "Move extrusion pricing to an LME-linked index with a collar so metal price risk is shared and not buried in unit price.", "Sourcing (commodity team)", "3-6 months", "Share of volume on index pricing", "All new extrusion contracts indexed"),
    ("R1", 4, "Pre-plan duty and customs treatment (bonded storage, duty-exemption eligibility) for imported extrusion and review quarterly against policy changes.", "Sourcing + Trade compliance", "0-6 months", "Duty paid versus plan", "No unplanned duty cost"),
    ("R2", 1, "Run side-by-side capability studies of Indian machine and tooling makers (Titan Engineering, Jyoti CNC, Bharat Forge) against incumbent equipment on one pilot line.", "Engineering + Supplier development", "0-9 months", "Pilot-line OEE and tolerance capability versus incumbent", "Within agreed gap of incumbent equipment"),
    ("R2", 2, "Dual-source critical equipment from non-China origins (for example Japan, Taiwan, Korea, Germany) and cut China's share of new equipment orders.", "Sourcing + Engineering", "3-12 months", "Share of new equipment value from non-China origin", "Majority of new orders non-China"),
    ("R2", 3, "Build a skilling and tool-maintenance programme in India (train and certify local technicians and process engineers) so tool upkeep no longer depends on imported engineers.", "Supplier development + HR partners", "0-12 months", "Certified Indian process engineers; mean time to repair", "Rising certified headcount; falling MTTR"),
    ("R2", 4, "Stock critical spares and consumable tooling in India with a defined buffer, and agree emergency supply terms with equipment makers.", "Supply chain planning", "0-3 months", "Days of critical-spares cover", "Cover for longest equipment lead time"),
    ("R3", 1, "Release volume to each new plant (Kanchipuram sites, Hosur, Kuppam) only on milestone gates for capacity, yield and quality.", "Sourcing + Operations", "0-12 months", "Gate passed on schedule", "No volume release without gate pass"),
    ("R3", 2, "Keep parallel supply across Tata, Yuzhan, Motherson and Aequs so one slip does not stop launch volume; cap any single plant's share.", "Sourcing (commodity team)", "0-6 months", "Largest single-plant share of volume", "Capped at an agreed ceiling"),
    ("R3", 3, "Run a weekly ramp review (schedule, OEE, first-pass yield) with each supplier and escalate any slip beyond an agreed number of weeks.", "Operations + Supplier quality", "0-12 months", "Schedule slip in weeks; first-pass yield versus plan", "Slip within tolerance; yield on plan"),
    ("R3", 4, "Keep contingency capacity at existing China or Vietnam enclosure lines during the first launch cycle of each new Indian line.", "Supply chain planning", "0-12 months", "Contingency capacity available", "Enough to cover the largest ramping plant"),
]
M0 = 5
for i, m in enumerate(mit):
    r_ = M0 + i
    put(mt, f"A{r_}", m[0], BOLD, align=CENTER, border=BOX)
    put(mt, f"B{r_}", f"=INDEX(Risks!$B${R0}:$B${R1},MATCH($A{r_},Risks!$A${R0}:$A${R1},0))", GREEN, align=WRAP, border=BOX)
    put(mt, f"C{r_}", m[1], BLACK, INT, align=CENTER, border=BOX)
    put(mt, f"D{r_}", m[2], BLACK, align=WRAP, border=BOX)
    put(mt, f"E{r_}", m[3], BLACK, align=WRAP, border=BOX)
    put(mt, f"F{r_}", m[4], BLACK, align=CENTER, border=BOX)
    put(mt, f"G{r_}", m[5], BLACK, align=WRAP, border=BOX)
    put(mt, f"H{r_}", m[6], BLACK, align=WRAP, border=BOX)
    mt.row_dimensions[r_].height = 50
M1 = M0 + len(mit) - 1
put(mt, f"A{M1 + 2}", "Residual risk after mitigation (from Risks sheet)", BOLD)
hdr(mt, M1 + 3, ["Risk", "Risk name", "", "Score before", "Score after", "Reduction", "", ""])
for i, rid in enumerate(["R1", "R2", "R3"]):
    r_ = M1 + 4 + i
    put(mt, f"A{r_}", rid, BOLD, align=CENTER, border=BOX)
    put(mt, f"B{r_}", f"=INDEX(Risks!$B${R0}:$B${R1},MATCH($A{r_},Risks!$A${R0}:$A${R1},0))", GREEN, align=WRAP, border=BOX)
    put(mt, f"D{r_}", f"=INDEX(Risks!$G${R0}:$G${R1},MATCH($A{r_},Risks!$A${R0}:$A${R1},0))", GREEN, INT, align=CENTER, border=BOX)
    put(mt, f"E{r_}", f"=INDEX(Risks!$M${R0}:$M${R1},MATCH($A{r_},Risks!$A${R0}:$A${R1},0))", GREEN, INT, align=CENTER, border=BOX)
    put(mt, f"F{r_}", f"=1-E{r_}/D{r_}", BLACK, PCT, align=CENTER, border=BOX)
    mt.row_dimensions[r_].height = 30
mt.freeze_panes = "A5"

# =========================================================== DASHBOARD
db = wb.create_sheet("Dashboard")
db.sheet_view.showGridLines = False
widths(db, {"A": 50, "B": 36, "C": 13, "D": 13, "E": 3})
put(db, "A1", "Dashboard: localization gap and supplier risk for mobile aluminium enclosures in India", TITLE)
put(db, "A2", "Every figure below is a formula driven by the ValueChain, Suppliers and Risks sheets. Scores are analyst judgments; see README and Sources.", SUB)
sec(db, 4, "Headline metrics", 4)
SR = f"Suppliers!$L${S0}:$L${S1}"
metrics = [
    (5, "Suppliers assessed", f"=COUNTA(Suppliers!$B${S0}:$B${S1})", INT),
    (6, "Suppliers rated High risk", f'=COUNTIF(Suppliers!$M${S0}:$M${S1},"High")', INT),
    (7, "Average supplier risk score (1 low, 5 high)", f"=AVERAGE({SR})", NUM2),
    (8, "Highest-risk supplier", f"=INDEX(Suppliers!$B${S0}:$B${S1},MATCH(MAX({SR}),{SR},0))", None),
    (9, "Value-chain stages assessed", f"=COUNTA(ValueChain!$B${V0}:$B${V1})", INT),
    (10, "Stages with import dependence of 4 or 5", f'=COUNTIF(ValueChain!$E${V0}:$E${V1},">=4")', INT),
    (11, "Stages with a Critical localization gap", f'=COUNTIF(ValueChain!$I${V0}:$I${V1},"Critical")', INT),
    (12, "Largest-gap stage", f"=INDEX(ValueChain!$B${V0}:$B${V1},MATCH(1,ValueChain!$H${V0}:$H${V1},0))", None),
]
for row, lab, fml, fmt in metrics:
    put(db, f"A{row}", lab, BLACK, border=BOX)
    put(db, f"B{row}", fml, BOLD, fmt, border=BOX, align=RIGHT)

sec(db, 14, "Top 3 risks (from the risk register)", 4)
hdr(db, 15, ["Risk", "Linked stage / supplier", "Score", "Rank"])
for k in range(1, 4):
    r_ = 15 + k
    m = f"MATCH({k},Risks!$H${R0}:$H${R1},0)"
    put(db, f"A{r_}", f"=INDEX(Risks!$B${R0}:$B${R1},{m})", GREEN, border=BOX, align=WRAP)
    put(db, f"B{r_}", f"=INDEX(Risks!$D${R0}:$D${R1},{m})", GREEN, border=BOX, align=WRAP)
    put(db, f"C{r_}", f"=INDEX(Risks!$G${R0}:$G${R1},{m})", GREEN, INT, border=BOX, align=CENTER)
    put(db, f"D{r_}", k, BLACK, INT, border=BOX, align=CENTER)

sec(db, 20, "Supplier risk ranking (highest risk first)", 4)
hdr(db, 21, ["Supplier", "Role in enclosure chain", "Risk score", "Rating"])
for k in range(1, len(sup) + 1):
    r_ = 21 + k
    m = f"MATCH({k},Suppliers!$N${S0}:$N${S1},0)"
    put(db, f"A{r_}", f'={k}&". "&INDEX(Suppliers!$B${S0}:$B${S1},{m})', GREEN, border=BOX)
    put(db, f"B{r_}", f"=INDEX(Suppliers!$C${S0}:$C${S1},{m})", GREEN, border=BOX, align=WRAP)
    put(db, f"C{r_}", f"=INDEX(Suppliers!$L${S0}:$L${S1},{m})", GREEN, NUM2, border=BOX, align=CENTER)
    put(db, f"D{r_}", f"=INDEX(Suppliers!$M${S0}:$M${S1},{m})", GREEN, border=BOX, align=CENTER)
    db.row_dimensions[r_].height = 27
DR0, DR1 = 22, 21 + len(sup)
db.conditional_formatting.add(f"D{DR0}:D{DR1}", FormulaRule(formula=[f'D{DR0}="High"'], fill=red))
db.conditional_formatting.add(f"D{DR0}:D{DR1}", FormulaRule(formula=[f'D{DR0}="Medium"'], fill=amb))
db.conditional_formatting.add(f"D{DR0}:D{DR1}", FormulaRule(formula=[f'D{DR0}="Low"'], fill=grn))

CX = DR1 + 2
sec(db, CX, "Context: sourced figures (inputs, not calculated)", 4)
hdr(db, CX + 1, ["Item", "Value", "Source", ""])
ctx = [
    ("ECMS applications approved to date (Aug 2026)", "106 (₹69,548 crore)", "[6]"),
    ("Mobile-enclosure investment in tranche 3 (Jan 2026)", "over ₹27,000 crore", "[17]"),
    ("Phone-enclosure extrusion imported today", "100%", "[4][8]"),
    ("Hindalco Kuppam extrusion plant: target start", "March 2027", "[7]"),
    ("Hindalco extrusion capacity with Kuppam + Silvassa", "109,000 t/yr (from 60,000)", "[8]"),
    ("Project 1: India local vs China import landed cost (base)", "+9.6% (India dearer)", "Project 1 model"),
    ("Project 1: India first-pass yield needed to match China", "about 95%", "Project 1 model"),
]
for i, (a, b, c) in enumerate(ctx):
    r_ = CX + 2 + i
    put(db, f"A{r_}", a, BLACK, border=BOX)
    put(db, f"B{r_}", b, BLUE, border=BOX, align=RIGHT)
    put(db, f"C{r_}", c, BLACK, border=BOX, align=CENTER)

ch1 = BarChart()
ch1.type = "bar"
ch1.style = 10
ch1.title = "Supplier risk score (1 = low, 5 = high)"
ch1.add_data(Reference(db, min_col=3, max_col=3, min_row=21, max_row=DR1), titles_from_data=True)
ch1.set_categories(Reference(db, min_col=1, min_row=DR0, max_row=DR1))
ch1.legend = None
ch1.gapWidth = 50
ch1.y_axis.scaling.min = 1
ch1.y_axis.scaling.max = 5
ch1.x_axis.scaling.orientation = "maxMin"
ch1.x_axis.tickLblPos = "low"
ch1.y_axis.delete = False
ch1.x_axis.delete = False
ch1.height, ch1.width = 11, 22
db.add_chart(ch1, "F4")

ch2 = BarChart()
ch2.type = "bar"
ch2.style = 10
ch2.title = "Value chain: import dependence vs India readiness (1-5)"
ch2.add_data(Reference(vc, min_col=5, max_col=6, min_row=6, max_row=V1), titles_from_data=True)
ch2.set_categories(Reference(vc, min_col=2, min_row=V0, max_row=V1))
ch2.gapWidth = 60
ch2.y_axis.scaling.min = 0
ch2.y_axis.scaling.max = 5
ch2.x_axis.scaling.orientation = "maxMin"
ch2.x_axis.tickLblPos = "low"
ch2.y_axis.delete = False
ch2.x_axis.delete = False
ch2.height, ch2.width = 13, 22
db.add_chart(ch2, "F27")

# =========================================================== SOURCES
so = wb.create_sheet("Sources")
so.sheet_view.showGridLines = False
widths(so, {"A": 6, "B": 56, "C": 70, "D": 90})
put(so, "A1", "Sources (retrieved 22 Oct 2026)", TITLE)
put(so, "A2", "Press and trade reports; some rely on unnamed sources. Verify plant status before relying on any single item.", SUB)
hdr(so, 4, ["Ref", "Report", "Used for", "URL"])
srcs = [
    (1, "Digitimes, Mar 2025: Jabil, Aequs, Tata export Apple components to China and Vietnam", "Motherson, Jabil, Aequs, Tata producing enclosures; exports", "https://www.digitimes.com/news/a20250303VL205/apple-india-mobile-components-manufacturing-supplier.html"),
    (2, "Digitimes, May 2025: Tata doubles Hosur iPhone enclosure production", "Tata output, Sept 2024 fire and restoration", "https://www.digitimes.com/news/a20250512VL204/iphone-apple-expansion-production-capacity.html"),
    (3, "Business Standard, Nov 2025: Apple steps up India push as suppliers scale", "TD Connex expansion, Aequs trial production, Yuzhan", "https://www.business-standard.com/companies/news/apple-india-suppliers-manufacturing-expansion-investments-supply-chain-125111500797_1.html"),
    (4, "EE Times India, Jan 2026: Government clears 22 more ECMS projects", "ECMS enclosure approvals; extrusion currently fully imported", "https://www.eetindia.co.in/government-clears-22-more-ecms-projects-deepening-indias-electronics-supply-chain/"),
    (5, "Deccan Herald, 2025: Govt approves 17 ECMS projects (tranche 2)", "Aequs laptop and smartwatch enclosures; Jabil", "https://www.deccanherald.com/amp/story/india%2Fgovt-approves-17-projects-with-rs-7172-cr-investment-for-electronics-component-manufacturing-3801087"),
    (6, "Ommcom News, 17 Aug 2026: Govt clears 31 ECMS projects", "106 approvals, ₹69,548 crore; Hosur, Kanchipuram enclosure plants", "https://ommcomnews.com/india-news/govt-clears-31-ecms-projects-worth-rs-7877-crore-106-approvals-now-total-rs-69548-crore/"),
    (7, "Steel News (Yieh), Sep 2025: Hindalco extrusion plant in Andhra Pradesh", "₹586 crore Kuppam plant, March 2027", "https://www.yieh.com/en/News/hindalco-to-build-aluminum-extrusion-plant-in-andhra-pradesh-for-iphones-chassis/156598"),
    (8, "Global SMT: India makes one in four iPhones; the aluminium body is still imported", "100% imported extrusion; Hindalco capacity", "https://www.globalsmt.net/?p=105209"),
    (9, "The420: Apple forms anchor vendor team under ECMS", "Tata supplies a limited share of enclosures", "https://the420.in/?p=61627"),
    (10, "Rest of World: Foxconn stops sending Chinese workers to India", "Held-up equipment; India lacked some technology", "https://restofworld.org/2025/china-foxconn-factoriesfoxconn-stops-sending-chinese-workers-to-india-iphone-factories/"),
    (11, "Deccan Herald / Bloomberg: Foxconn recall of Chinese staff", "About 300 engineers recalled; Yuzhan makes older-iPhone enclosures", "https://www.deccanherald.com/amp/story/business%2Ffoxconns-recall-of-more-chinese-staff-tests-apples-india-push-3694305"),
    (12, "NewsBytes, Sep 2025: Apple strengthens India manufacturing with machinery", "About 35 machinery makers; TEAL, Jyoti CNC, Bharat Forge, Wipro", "https://www.newsbytesapp.com/news/business/apple-boosts-iphone-production-in-india-amid-china-tensions/story"),
    (13, "Storyboard18, 26 Aug 2026: India pushes Apple suppliers to deepen component manufacturing", "Local firms export mechanicals and housings to China and Vietnam", "https://www.storyboard18.com/brand-marketing/india-pushes-apple-suppliers-to-deepen-local-component-manufacturing-report-108791.htm"),
    (14, "Reuters via AOL: Jabil to set up South India plant ($238M)", "Jabil Trichy plant", "https://www.aol.com/news/apple-supplier-jabil-set-south-051928977.html"),
    (15, "CB Insights: TDConnex profile", "TD Connex capabilities (CNC, injection, stamping, LSR)", "https://www.cbinsights.com/company/tdconnex"),
    (17, "Legal Parivar News: ₹41,863 crore ECMS projects, focus on mobile enclosures", "Over ₹27,000 crore enclosure investment", "https://news.legalparivar.com/india-clears-%E2%82%B941863-cr-ecms-projects-focus-on-mobile-enclosures/"),
    (18, "Gulf News / Bloomberg: Tata to add up to 45,000 workers at Hosur", "Hiring scale; competing housing suppliers", "https://gulfnews.com/business/retail/indias-tata-to-add-up-to-45000-workers-at-iphone-parts-plant-1.1667294056340"),
    (19, "Gulf News: Foxconn making iPhone casings locally for the first time", "Foxconn entry; Tata previously the only local enclosure maker", "https://gulfnews.com/amp/story/technology%2Ffoxconn-expands-in-india-now-making-iphone-casings-locally-for-the-first-time-says-report-1.500170272"),
    (20, "The Tribune / ANI, 2 Jan 2026: Third tranche of ECMS", "Motherson Tamil Nadu project, 5,741 jobs", "https://www.tribuneindia.com/news/business/third-tranche-of-ecms-approved-with-rs-41863-crore-investment-covering-22-applications/"),
    (21, "Communications Today: Apple ramps up India value addition", "Jabil plastic AirPods enclosures; Aequs non-iPhone enclosures", "https://www.communicationstoday.co.in/apple-is-ramping-up-india-value-addition-with-flextronics-jabil-aequs-motherson/"),
    (22, "AlCircle / Business Standard: Hindalco investment plans", "Hindalco $10bn five-year plan including Dahej smelter", "https://www.business-standard.com/amp/topic/hindalco"),
    (23, "Deccan Herald / Bloomberg (as ref 11)", "Beijing curbing technology and equipment transfers to India", "https://www.deccanherald.com/amp/story/business%2Ffoxconns-recall-of-more-chinese-staff-tests-apples-india-push-3694305"),
    (24, "India Briefing: How Apple made India a manufacturing hub", "Apple supplier list incl. Titan Engineering & Automation, VVDN", "https://www.india-briefing.com/news/apple-india-local-manufacturing-growth-40070.html/"),
    (25, "AlCircle / Procurement Resource, Oct 2026: LME aluminium falls to $3,107/t", "Aluminium price volatility", "https://www.procurementresource.com/news-and-articles/lme-aluminium-falls-to-3-107-dollars"),
    (26, "J2T: Aluminium price forecast 2026-2030", "Analyst range for 2026 average price", "https://j2t.com/solutions/blogview/aluminium-price-forecast/"),
]
for i, s in enumerate(srcs):
    r_ = 5 + i
    for j, v in enumerate(s):
        put(so, f"{L(1 + j)}{r_}", v, BLACK, align=WRAP, border=BOX)
    so.row_dimensions[r_].height = 34

for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

order = ["README", "Dashboard", "ValueChain", "Suppliers", "Risks", "Mitigation", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("saved", OUT)
