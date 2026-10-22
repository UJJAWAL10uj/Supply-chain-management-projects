"""
Build the India vs China vs Vietnam TCO & scenario model for a CNC-machined
aluminium smartphone enclosure delivered to an assembly plant near Chennai.

Every calculated number in the workbook is an Excel formula; this script only
lays out inputs, labels and formulas. Run:  python build_model.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference

OUT = "/mnt/user-data/outputs/Apple_Sourcing_TCO_Model_Aluminium_Enclosure.xlsx"

FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
WHITE_B = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=10, italic=True, color="595959")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

USD3 = '$#,##0.000;($#,##0.000);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT1 = '0.0%;(0.0%);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM0 = '#,##0;(#,##0);"-"'
MULT = '0.00"x"'

wb = Workbook()


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        put(ws, f"{L(start_col + i)}{row}", lab, WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)


def section(ws, row, text, ncols):
    for i in range(1, ncols + 1):
        ws[f"{L(i)}{row}"].fill = SEC_FILL
    put(ws, f"A{row}", text, BOLD, fill=SEC_FILL)


# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 110
put(ws, "B2", "TCO & Scenario Model: Aluminium Smartphone Enclosure, India vs China vs Vietnam", TITLE)
put(ws, "B3", "Independent academic project built from public data and stated assumptions. Not Apple data.", SUB)
rows = [
    ("Question", "Where should a CNC-machined aluminium enclosure be sourced from to be delivered to a final-assembly plant near Chennai: a local Indian supplier, or imports from China or Vietnam? And how robust is the answer to tariff, freight, volume, metal-price, incentive and yield shocks?"),
    ("Part modelled", "Unibody-style CNC-machined 6xxx aluminium enclosure, 28 g finished, 85 g blank, anodised. Delivered cost to the assembly plant (Chennai area), USD per unit, at 5 million units per year per supplier line."),
    ("Cost layers", "Net material (after chip-scrap credit) > machining > direct labour > finishing > yield loss > fixed cost per unit (tooling + overhead) > supplier margin > ECMS incentive (India only) > freight > insurance > Basic Customs Duty + Social Welfare Surcharge + other duties > clearance and haulage > pipeline carrying cost."),
    ("Sheets", "Inputs (all assumptions) > Scenarios (11 shock definitions) > Calc (33-column engine, same formula in every column) > Summary (results, cost stack, charts) > Sensitivity (two-way grids and break-evens) > Incoterms (FOB vs CIF vs DDP, who pays what) > Sources (what is sourced vs assumed)."),
    ("How to use", "Change only BLUE cells on Inputs or Scenarios. Everything else recalculates. Yellow cells mark the key levers. To test your own idea, overwrite a scenario row on the Scenarios sheet."),
    ("Colour code", "Blue = hard-coded input. Black = formula. Green = link to another sheet. Yellow fill = key assumption or lever."),
    ("Important limitation", "Process inputs (cycle time, machine rate, labour rate, yield, tooling, freight rates) are illustrative values in plausible ranges, NOT quotes. Public anchors (LME price, ECMS incentive range, tariff regime) are listed on the Sources sheet. Replace the blue cells with RFQ or benchmark data before drawing real conclusions."),
    ("Simplifications", "IGST/GST is excluded because it is creditable to a registered manufacturer. Scrapped parts are not credited with metal value (conservative). Anti-dumping duty is an input set to zero. Vietnam BCD of 5% is a placeholder for a possible ASEAN-India FTA preferential rate and must be verified against HS code and rules of origin. US tariffs on the finished device are discussed in the memo, not modelled here."),
]
r = 5
for k, v in rows:
    put(ws, f"B{r}", k, BOLD, align=WRAP)
    put(ws, f"C{r}", v, BLACK, align=WRAP)
    ws.row_dimensions[r].height = 62 if len(v) > 220 else 46
    r += 1

# ---------------------------------------------------------------- INPUTS
wi = wb.create_sheet("Inputs")
wi.sheet_view.showGridLines = False
for col, w in zip("ABCDEF", [46, 14, 14, 14, 14, 78]):
    wi.column_dimensions[col].width = w
put(wi, "A1", "Inputs and assumptions", TITLE)
put(wi, "A2", "Blue = input you can change. Yellow = key lever. All USD unless stated.", SUB)

section(wi, 4, "Global parameters", 6)
header_row(wi, 5, ["Parameter", "Value", "Unit", "", "", "Basis / source"])
glob = [
    # row, label, value, unit, fmt, note, key
    (6, "Finished part weight", 28, "g", NUM0, "Assumption: typical smartphone-class machined enclosure", False),
    (7, "Blank (billet/extrusion) weight", 85, "g", NUM0, "Assumption: ~3:1 buy-to-fly ratio typical for CNC unibody parts", False),
    (8, "Chip-scrap recovery value (% of LME)", 0.60, "%", PCT1, "Assumption: machining chips resold at a discount to LME", False),
    (9, "LME aluminium cash price", 3134, "USD/t", NUM0, "Source: AlCircle, LME cash offer USD 3,134.5/t on 6 Oct 2026 (see Sources)", True),
    (10, "Annual volume per supplier line (base)", 5000000, "units/yr", NUM0, "Assumption: single dedicated line", True),
    (11, "Cost of capital (annual)", 0.10, "%", PCT1, "Assumption: used for pipeline inventory carrying cost", False),
    (12, "Tooling amortisation period", 3, "years", NUM0, "Assumption: product life of one enclosure generation", False),
    (13, "Packed chargeable weight per unit", 0.07, "kg", NUM2, "Assumption: part plus protective packaging", False),
    (14, "Social Welfare Surcharge (% of BCD)", 0.10, "%", PCT1, "India customs: SWS is 10% of BCD (verify current rate)", False),
    (15, "Seller markup on freight + insurance under CIF", 0.08, "%", PCT1, "Assumption: seller margin on arranging carriage", True),
    (16, "Seller markup on duty, clearance and carrying under DDP", 0.06, "%", PCT1, "Assumption: seller margin on carrying duty and clearance risk", True),
]
for row, lab, val, unit, fmt, note, key in glob:
    put(wi, f"A{row}", lab, BLACK, border=BOX)
    put(wi, f"B{row}", val, BLUE, fmt, KEY_FILL if key else None, border=BOX)
    put(wi, f"C{row}", unit, BLACK, border=BOX)
    put(wi, f"F{row}", note, BLACK, align=WRAP)

section(wi, 18 - 1, "Country parameters (delivered to assembly plant near Chennai)", 6)
header_row(wi, 18, ["Parameter", "Unit", "India", "China", "Vietnam", "Basis / source"])
ctry = [
    (19, "Imported into India plant? (1 = yes, 0 = local)", "flag", [0, 1, 1], NUM0, False, "Structural: India supplier is local, others ship in"),
    (20, "Regional alloy / fabrication premium over LME", "USD/t", [520, 330, 450], NUM0, False, "Assumption: India includes duty on imported primary metal"),
    (21, "CNC cycle time", "min/unit", [9.0, 8.0, 8.8], NUM2, True, "Assumption: India slightly slower on learning curve"),
    (22, "Machine-hour rate (depreciation, power, tooling)", "USD/hr", [14.0, 11.5, 12.5], NUM2, True, "Assumption: India higher capital cost and import duty on machines"),
    (23, "Operators per machine", "ratio", [0.25, 0.25, 0.25], NUM2, False, "Assumption: 1 operator tends 4 machines"),
    (24, "Direct labour rate", "USD/hr", [3.2, 6.5, 3.8], NUM2, False, "Assumption: loaded hourly cost, industrial region"),
    (25, "Anodising / finishing cost", "USD/unit", [0.55, 0.40, 0.48], USD2, False, "Assumption"),
    (26, "First-pass yield (base)", "%", [0.86, 0.93, 0.90], PCT1, True, "Assumption: India ramp-up yield below mature China line"),
    (27, "Tooling and fixtures capex", "USD", [1800000, 1400000, 1600000], USD0, False, "Assumption"),
    (28, "Fixed overhead", "USD/yr", [1200000, 1500000, 1100000], USD0, False, "Assumption: rent, supervision, QA"),
    (29, "Supplier margin on cost", "%", [0.09, 0.08, 0.09], PCT1, False, "Assumption"),
    (30, "Sea / road freight rate", "USD/kg", [0.12, 0.45, 0.55], NUM2, False, "Assumption: India = trucking to plant"),
    (31, "Air freight rate", "USD/kg", [0, 4.2, 4.8], NUM2, False, "Assumption"),
    (32, "Air share of shipments", "%", [0, 0.40, 0.40], PCT1, True, "Assumption: part of volume flown to protect launch schedules"),
    (33, "Transit days by sea / road", "days", [2, 25, 18], NUM0, False, "Assumption"),
    (34, "Transit days by air", "days", [2, 4, 4], NUM0, False, "Assumption"),
    (35, "Cargo insurance", "% of price", [0.001, 0.003, 0.003], '0.00%', False, "Assumption"),
    (36, "Basic Customs Duty (BCD) into India", "%", [0, 0.10, 0.05], PCT1, True, "10% on metal mechanical phone parts per Jan 2024 notice; VERIFY current CBIC rate. Vietnam 5% = placeholder for FTA rate"),
    (37, "Additional duties (anti-dumping / safeguard)", "% of CIF", [0, 0, 0], PCT1, False, "Input kept at zero; use to test trade-remedy risk"),
    (38, "Customs clearance and handling", "% of CIF", [0, 0.004, 0.004], '0.00%', False, "Assumption"),
    (39, "Port-to-plant haulage", "USD/kg", [0, 0.05, 0.05], NUM2, False, "Assumption: Chennai port is close to the plant"),
    (40, "ECMS turnover-linked incentive rate", "% of sales", [0.07, 0, 0], PCT1, True, "ECMS enclosures: 6-8% initially, falling to 3-5% (Vision IAS summary). 7% = midpoint"),
    (41, "Share of incentive passed to buyer", "%", [0.50, 0, 0], PCT1, True, "Assumption: negotiated; the rest stays with the supplier"),
]
for row, lab, unit, vals, fmt, key, note in ctry:
    put(wi, f"A{row}", lab, BLACK, border=BOX)
    put(wi, f"B{row}", unit, BLACK, border=BOX)
    for j, v in enumerate(vals):
        put(wi, f"{L(3 + j)}{row}", v, BLUE, fmt, KEY_FILL if key else None, border=BOX)
    put(wi, f"F{row}", note, BLACK, align=WRAP)
wi.freeze_panes = "A6"

# ---------------------------------------------------------------- SCENARIOS
wsn = wb.create_sheet("Scenarios")
wsn.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [52, 15, 13, 13, 15, 17, 15, 70]):
    wsn.column_dimensions[col].width = w
put(wsn, "A1", "Scenario levers", TITLE)
put(wsn, "A2", "Each row is one scenario. Multipliers: 1.00 = no change. BCD change is in percentage points added to the import duty rate (floored at 0%). Blue cells can be edited.", SUB)
header_row(wsn, 4, ["Scenario", "Import duty change (pts)", "Freight multiplier", "Volume multiplier", "Aluminium price multiplier", "ECMS pass-through multiplier", "India yield change (pts)", "What it represents"])
wsn.row_dimensions[4].height = 42
scen = [
    ("Base case", 0, 1, 1, 1, 1, 0, "Current duty, freight, metal price and incentive terms"),
    ("Import duty +10 pts", 0.10, 1, 1, 1, 1, 0, "India raises BCD on imported enclosures to protect local capacity"),
    ("Import duty exempt (BCD to 0)", -1.0, 1, 1, 1, 1, 0, "Duty cut to zero on imported parts (e.g. trade deal or exemption)"),
    ("Freight spike x2.5", 0, 2.5, 1, 1, 1, 0, "Route disruption or air-freight capacity squeeze"),
    ("Volume -10%", 0, 1, 0.90, 1, 1, 0, "Weaker demand; fixed cost spread over fewer units"),
    ("Volume +10%", 0, 1, 1.10, 1, 1, 0, "Stronger demand; better fixed-cost absorption"),
    ("Aluminium +15%", 0, 1, 1, 1.15, 1, 0, "Metal price rally"),
    ("ECMS delayed (0% pass-through)", 0, 1, 1, 1, 0, 0, "Incentive delayed or not passed to the buyer"),
    ("India yield -4 pts", 0, 1, 1, 1, 1, -0.04, "Slower yield ramp at the Indian supplier"),
    ("Import-shock stack (duty +10, freight x2.5, Al +15%)", 0.10, 2.5, 1, 1.15, 1, 0, "Everything goes wrong for importers at once"),
    ("India-risk stack (vol -10%, ECMS 0%, yield -4 pts)", 0, 1, 0.90, 1, 0, -0.04, "Everything goes wrong for the local supplier at once"),
]
for i, s in enumerate(scen):
    row = 5 + i
    put(wsn, f"A{row}", s[0], BLUE, border=BOX)
    fmts = [PCT1, MULT, MULT, MULT, MULT, PCT1]
    for j in range(6):
        put(wsn, f"{L(2 + j)}{row}", s[1 + j], BLUE, fmts[j], KEY_FILL, border=BOX)
    put(wsn, f"H{row}", s[7], BLACK, align=WRAP)
NS = len(scen)
S_FIRST, S_LAST = 5, 5 + NS - 1  # 5..15

# ---------------------------------------------------------------- CALC
wc = wb.create_sheet("Calc")
wc.sheet_view.showGridLines = False
wc.column_dimensions["A"].width = 46
wc.column_dimensions["B"].width = 12
put(wc, "A1", "Calculation engine: one column per scenario x country (identical formulas in every column)", TITLE)
put(wc, "A2", "Do not edit. Inputs come from the Inputs and Scenarios sheets. USD per unit unless stated.", SUB)

CALC_FIRST = 3
CALC_LAST = CALC_FIRST + NS * 3 - 1  # 35 -> AI
cF, cL = L(CALC_FIRST), L(CALC_LAST)

labels = {
    4: ("Scenario", ""), 5: ("Country", ""), 6: ("Lookup key", ""),
    7: ("Import duty change", "pts"), 8: ("Freight multiplier", "x"), 9: ("Volume multiplier", "x"),
    10: ("Aluminium price multiplier", "x"), 11: ("ECMS pass-through multiplier", "x"), 12: ("India yield change", "pts"),
    13: ("COUNTRY INPUTS (looked up from Inputs)", ""),
    14: ("Imported into India plant? (1/0)", "flag"), 15: ("Alloy / fabrication premium", "USD/t"),
    16: ("CNC cycle time", "min"), 17: ("Machine-hour rate", "USD/hr"), 18: ("Operators per machine", "ratio"),
    19: ("Direct labour rate", "USD/hr"), 20: ("Finishing cost", "USD"), 21: ("First-pass yield (base)", "%"),
    22: ("Tooling capex", "USD"), 23: ("Fixed overhead", "USD/yr"), 24: ("Supplier margin", "%"),
    25: ("Sea / road rate", "USD/kg"), 26: ("Air rate", "USD/kg"), 27: ("Air share", "%"),
    28: ("Transit days sea / road", "days"), 29: ("Transit days air", "days"), 30: ("Cargo insurance rate", "%"),
    31: ("BCD base rate", "%"), 32: ("Additional duty rate", "%"), 33: ("Clearance rate", "%"),
    34: ("Haulage", "USD/kg"), 35: ("ECMS incentive rate", "%"), 36: ("ECMS pass-through (base)", "%"),
    37: ("COST BUILD-UP (USD per unit)", ""),
    38: ("Aluminium price incl. premium", "USD/kg"), 39: ("Gross material", "USD"), 40: ("Chip-scrap credit", "USD"),
    41: ("Net material", "USD"), 42: ("Machining", "USD"), 43: ("Direct labour", "USD"), 44: ("Finishing", "USD"),
    45: ("Conversion subtotal", "USD"), 46: ("Effective first-pass yield", "%"), 47: ("Yield-loss cost", "USD"),
    48: ("Fixed cost per unit (tooling + overhead)", "USD"), 49: ("Total supplier cost", "USD"),
    50: ("Supplier margin", "USD"), 51: ("Ex-works price before incentive", "USD"),
    52: ("ECMS incentive passed to buyer", "USD"), 53: ("Net ex-works price", "USD"),
    54: ("LOGISTICS, DUTIES AND CARRYING COST", ""),
    55: ("Blended freight rate", "USD/kg"), 56: ("International freight", "USD"), 57: ("Cargo insurance", "USD"),
    58: ("CIF / delivered value", "USD"), 59: ("Effective BCD rate", "%"), 60: ("Basic Customs Duty", "USD"),
    61: ("Social Welfare Surcharge", "USD"), 62: ("Additional duties", "USD"), 63: ("Customs clearance and handling", "USD"),
    64: ("Port-to-plant haulage", "USD"), 65: ("Pipeline carrying cost", "USD"),
    66: ("TOTAL LANDED COST PER UNIT", "USD"), 67: ("Landed cost excluding ECMS benefit", "USD"),
    68: ("Annual landed spend at scenario volume", "USD m"),
}
for r_, (lab, unit) in labels.items():
    is_sec = lab.isupper() and unit == "" and r_ in (13, 37, 54)
    put(wc, f"A{r_}", lab, BOLD if (is_sec or r_ in (66,)) else BLACK, fill=SEC_FILL if is_sec else (TOT_FILL if r_ == 66 else None))
    put(wc, f"B{r_}", unit, SUB, fill=SEC_FILL if is_sec else (TOT_FILL if r_ == 66 else None))
    if is_sec:
        for ci in range(CALC_FIRST, CALC_LAST + 1):
            wc[f"{L(ci)}{r_}"].fill = SEC_FILL

# map: calc row -> Inputs country row
in_rows = {14: 19, 15: 20, 16: 21, 17: 22, 18: 23, 19: 24, 20: 25, 21: 26, 22: 27, 23: 28, 24: 29,
           25: 30, 26: 31, 27: 32, 28: 33, 29: 34, 30: 35, 31: 36, 32: 37, 33: 38, 34: 39, 35: 40, 36: 41}
in_fmt = {14: NUM0, 15: NUM0, 16: NUM2, 17: NUM2, 18: NUM2, 19: NUM2, 20: USD2, 21: PCT1, 22: USD0, 23: USD0,
          24: PCT1, 25: NUM2, 26: NUM2, 27: PCT1, 28: NUM0, 29: NUM0, 30: '0.00%', 31: PCT1, 32: PCT1,
          33: '0.00%', 34: NUM2, 35: PCT1, 36: PCT1}
lever_cols = {7: "B", 8: "C", 9: "D", 10: "E", 11: "F", 12: "G"}
lever_fmt = {7: PCT1, 8: MULT, 9: MULT, 10: MULT, 11: MULT, 12: PCT1}

for i in range(NS):
    for j in range(3):
        ci = CALC_FIRST + i * 3 + j
        c = L(ci)
        wc.column_dimensions[c].width = 15
        # scenario / country / key
        put(wc, f"{c}4", f"=Scenarios!$A${S_FIRST + i}", GREEN, fill=None, align=CENTER)
        put(wc, f"{c}5", f"=Inputs!${L(3 + j)}$18", GREEN, align=CENTER)
        put(wc, f"{c}6", f'={c}4&"|"&{c}5', BLACK, align=CENTER)
        wc.row_dimensions[4].height = 54
        # levers
        for lr, scol in lever_cols.items():
            put(wc, f"{c}{lr}",
                f"=INDEX(Scenarios!${scol}${S_FIRST}:${scol}${S_LAST},MATCH({c}$4,Scenarios!$A${S_FIRST}:$A${S_LAST},0))",
                GREEN, lever_fmt[lr])
        # country inputs
        for cr, ir in in_rows.items():
            put(wc, f"{c}{cr}", f"=INDEX(Inputs!$C${ir}:$E${ir},MATCH({c}$5,Inputs!$C$18:$E$18,0))", GREEN, in_fmt[cr])
        f = {}
        f[38] = f"=(Inputs!$B$9*{c}$10+{c}15)/1000"
        f[39] = f"=Inputs!$B$7/1000*{c}38"
        f[40] = f"=-(Inputs!$B$7-Inputs!$B$6)/1000*(Inputs!$B$9*{c}$10/1000)*Inputs!$B$8"
        f[41] = f"={c}39+{c}40"
        f[42] = f"={c}16/60*{c}17"
        f[43] = f"={c}16/60*{c}18*{c}19"
        f[44] = f"={c}20"
        f[45] = f"={c}42+{c}43+{c}44"
        f[46] = f"=MIN(1,{c}21+IF({c}$5=Inputs!$C$18,{c}$12,0))"
        f[47] = f"=({c}41+{c}45)*(1/{c}46-1)"
        f[48] = f"=({c}22/Inputs!$B$12+{c}23)/(Inputs!$B$10*{c}$9)"
        f[49] = f"={c}41+{c}45+{c}47+{c}48"
        f[50] = f"={c}49*{c}24"
        f[51] = f"={c}49+{c}50"
        f[52] = f"=-{c}35*{c}51*MIN(1,{c}36*{c}$11)"
        f[53] = f"={c}51+{c}52"
        f[55] = f"={c}27*{c}26+(1-{c}27)*{c}25"
        f[56] = f"=Inputs!$B$13*{c}55*{c}$8"
        f[57] = f"={c}30*{c}53"
        f[58] = f"={c}53+{c}56+{c}57"
        f[59] = f"={c}14*MAX(0,{c}31+{c}$7)"
        f[60] = f"={c}59*{c}58"
        f[61] = f"=Inputs!$B$14*{c}60"
        f[62] = f"={c}14*{c}32*{c}58"
        f[63] = f"={c}14*{c}33*{c}58"
        f[64] = f"={c}14*{c}34*Inputs!$B$13"
        f[65] = f"={c}53*Inputs!$B$11*({c}27*{c}29+(1-{c}27)*{c}28)/365"
        f[66] = f"={c}53+{c}56+{c}57+{c}60+{c}61+{c}62+{c}63+{c}64+{c}65"
        f[67] = f"={c}66-{c}52"
        f[68] = f"={c}66*Inputs!$B$10*{c}$9/1000000"
        fm = {38: USD3, 39: USD3, 40: USD3, 41: USD3, 42: USD3, 43: USD3, 44: USD3, 45: USD3, 46: PCT1, 47: USD3,
              48: USD3, 49: USD3, 50: USD3, 51: USD3, 52: USD3, 53: USD3, 55: NUM2, 56: USD3, 57: USD3, 58: USD3,
              59: PCT1, 60: USD3, 61: USD3, 62: USD3, 63: USD3, 64: USD3, 65: USD3, 66: USD3, 67: USD3, 68: '#,##0.0'}
        for rr, ff in f.items():
            put(wc, f"{c}{rr}", ff, BOLD if rr == 66 else BLACK, fm[rr], TOT_FILL if rr == 66 else None)
wc.freeze_panes = "C7"

KEYS = f"Calc!${cF}$6:${cL}$6"


def calc_pick(row, key_expr):
    return f"INDEX(Calc!${cF}${row}:${cL}${row},MATCH({key_expr},{KEYS},0))"


# ---------------------------------------------------------------- SUMMARY
wsu = wb.create_sheet("Summary")
wsu.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [54, 14, 14, 14, 16, 16, 16, 3]):
    wsu.column_dimensions[col].width = w
put(wsu, "A1", "Results: landed cost of an aluminium enclosure delivered near Chennai (USD per unit)", TITLE)
put(wsu, "A2", "All values are formulas driven by the Inputs and Scenarios sheets. Green cell = cheapest source in that scenario.", SUB)
section(wsu, 4, "Landed cost per unit by scenario", 7)
put(wsu, "B4", "Local", BOLD, fill=SEC_FILL, align=CENTER)
put(wsu, "C4", "Import", BOLD, fill=SEC_FILL, align=CENTER)
put(wsu, "D4", "Import", BOLD, fill=SEC_FILL, align=CENTER)
put(wsu, "A5", "Scenario", WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
for j in range(3):
    put(wsu, f"{L(2 + j)}5", f"=Inputs!${L(3 + j)}$18", WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
for col, lab in zip("EFG", ["Cheapest source", "India vs China (%)", "India vs Vietnam (%)"]):
    put(wsu, f"{col}5", lab, WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
wsu.row_dimensions[5].height = 32
for i in range(NS):
    r_ = 6 + i
    put(wsu, f"A{r_}", f"=Scenarios!$A${S_FIRST + i}", GREEN, border=BOX)
    for j in range(3):
        col = L(2 + j)
        put(wsu, f"{col}{r_}", f'={calc_pick(66, f"$A{r_}&\"|\"&{col}$5")}', BLACK, USD3, border=BOX)
    put(wsu, f"E{r_}", f"=INDEX($B$5:$D$5,MATCH(MIN(B{r_}:D{r_}),B{r_}:D{r_},0))", BOLD, border=BOX, align=CENTER)
    put(wsu, f"F{r_}", f"=B{r_}/C{r_}-1", BLACK, PCT1, border=BOX)
    put(wsu, f"G{r_}", f"=B{r_}/D{r_}-1", BLACK, PCT1, border=BOX)
SC_LAST = 6 + NS - 1  # 16
wsu.conditional_formatting.add(
    f"B6:D{SC_LAST}",
    FormulaRule(formula=["B6=MIN($B6:$D6)"], fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")),
)

section(wsu, 18, "Base-case headline metrics", 7)
base_key = lambda ctry_ref: f'Summary!$A$6&"|"&{ctry_ref}'
heads = [
    (19, "Cheapest source (base case)", "=E6", None),
    (20, "India local landed cost (USD/unit)", "=B6", USD3),
    (21, "China imported landed cost (USD/unit)", "=C6", USD3),
    (22, "India premium vs China (%)", "=F6", PCT1),
    (23, "Break-even import duty (BCD) at which India local = China", "=Sensitivity!B26", PCT1),
    (24, "Break-even freight multiplier at which India local = China", "=Sensitivity!B27", MULT),
    (25, "Annual spend gap, India minus China, at base volume (USD m)", "=(B6-C6)*Inputs!$B$10/1000000", '#,##0.00;(#,##0.00);"-"'),
]
for row, lab, fml, fmt in heads:
    put(wsu, f"A{row}", lab, BLACK, border=BOX)
    put(wsu, f"B{row}", fml, BOLD, fmt, border=BOX, align=Alignment(horizontal="right"))

section(wsu, 28, "Base-case cost stack (USD per unit)", 7)
put(wsu, "A29", "Component", WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
for j in range(3):
    put(wsu, f"{L(2 + j)}29", f"=Inputs!${L(3 + j)}$18", WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
stack = [
    ("Net material", [41]), ("Machining", [42]), ("Direct labour", [43]), ("Finishing", [44]),
    ("Yield loss", [47]), ("Fixed cost (tooling + overhead)", [48]), ("Supplier margin", [50]),
    ("ECMS incentive to buyer", [52]), ("International freight", [56]), ("Cargo insurance", [57]),
    ("Import duties (BCD + SWS + other)", [60, 61, 62]), ("Clearance and haulage", [63, 64]),
    ("Pipeline carrying cost", [65]),
]
for k, (lab, rws) in enumerate(stack):
    r_ = 30 + k
    put(wsu, f"A{r_}", lab, BLACK, border=BOX)
    for j in range(3):
        col = L(2 + j)
        parts = "+".join(calc_pick(rr, base_key(f"{col}$29")) for rr in rws)
        put(wsu, f"{col}{r_}", f"={parts}", BLACK, USD3, border=BOX)
ST_LAST = 30 + len(stack) - 1  # 42
put(wsu, f"A{ST_LAST + 1}", "Total landed cost", BOLD, fill=TOT_FILL, border=BOX)
put(wsu, f"A{ST_LAST + 2}", "Check vs Calc (should be 0)", SUB)
for j in range(3):
    col = L(2 + j)
    put(wsu, f"{col}{ST_LAST + 1}", f"=SUM({col}30:{col}{ST_LAST})", BOLD, USD3, TOT_FILL, border=BOX)
    put(wsu, f"{col}{ST_LAST + 2}", f"=ROUND({col}{ST_LAST + 1}-{calc_pick(66, base_key(f'{col}$29'))},6)", SUB, '0.000000')

ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = "India local vs China imported: landed-cost gap (%)  [+ = India costlier]"
ch.add_data(Reference(wsu, min_col=6, max_col=6, min_row=5, max_row=SC_LAST), titles_from_data=True)
ch.set_categories(Reference(wsu, min_col=1, min_row=6, max_row=SC_LAST))
ch.height, ch.width = 12, 24
ch.gapWidth = 50
ch.legend = None
ch.y_axis.number_format = '0%'
ch.x_axis.scaling.orientation = "maxMin"
ch.x_axis.tickLblPos = "low"
ch.y_axis.delete = False
ch.x_axis.delete = False
wsu.add_chart(ch, "I4")

ch2 = BarChart()
ch2.type = "col"
ch2.grouping = "stacked"
ch2.overlap = 100
ch2.title = "Base-case cost stack (USD/unit)"
ch2.add_data(Reference(wsu, min_col=1, max_col=4, min_row=30, max_row=ST_LAST), from_rows=True, titles_from_data=True)
ch2.set_categories(Reference(wsu, min_col=2, max_col=4, min_row=29))
ch2.height, ch2.width = 12, 22
ch2.y_axis.delete = False
ch2.x_axis.delete = False
wsu.add_chart(ch2, "I30")

# ---------------------------------------------------------------- SENSITIVITY
wss = wb.create_sheet("Sensitivity")
wss.sheet_view.showGridLines = False
wss.column_dimensions["A"].width = 58
for col in "BCDEFGH":
    wss.column_dimensions[col].width = 13
put(wss, "A1", "Sensitivity grids and break-even points", TITLE)
put(wss, "A2", "India advantage = China imported landed cost minus India local landed cost. Positive (green) = India local is cheaper.", SUB)

section(wss, 4, "Reference values pulled from Calc (base case)", 8)
cn = "Inputs!$D$18"
ind = "Inputs!$C$18"
refs = [
    (5, "China: net ex-works price", calc_pick(53, base_key(cn)), USD3),
    (6, "China: cargo insurance", calc_pick(57, base_key(cn)), USD3),
    (7, "China: international freight (base)", calc_pick(56, base_key(cn)), USD3),
    (8, "China: port-to-plant haulage", calc_pick(64, base_key(cn)), USD3),
    (9, "China: pipeline carrying cost", calc_pick(65, base_key(cn)), USD3),
    (10, "Social Welfare Surcharge (% of BCD)", "Inputs!$B$14", PCT1),
    (11, "China: additional duty rate", calc_pick(32, base_key(cn)), PCT1),
    (12, "China: clearance rate", calc_pick(33, base_key(cn)), '0.00%'),
    (13, "China: base BCD rate", calc_pick(59, base_key(cn)), PCT1),
    (14, "India local: landed cost (base)", calc_pick(66, base_key(ind)), USD3),
    (15, "China imported: landed cost (base)", calc_pick(66, base_key(cn)), USD3),
]
for row, lab, fml, fmt in refs:
    put(wss, f"A{row}", lab, BLACK, border=BOX)
    put(wss, f"B{row}", f"={fml}", GREEN, fmt, border=BOX)


def china_landed(r, m):
    k = f"((1+$B$10)*{r}+$B$11+$B$12)"
    return f"(($B$5+$B$6+{m}*$B$7)*(1+{k})+$B$8+$B$9)"


# Grid A
section(wss, 17, "Grid A: India advantage (USD/unit) vs import duty rate (rows) and freight multiplier (columns)", 8)
put(wss, "A18", "BCD on China imports (rows) / Freight multiplier (columns)", BOLD)
fm_list = [1.0, 1.5, 2.0, 2.5, 3.0, 4.0]
bcd_list = [0.0, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30]
for j, m in enumerate(fm_list):
    put(wss, f"{L(2 + j)}18", m, BLUE, MULT, KEY_FILL, CENTER, BOX)
# Break-evens (rows 26-27) placed after grid A rows 19..25
for i, rr in enumerate(bcd_list):
    row = 19 + i
    put(wss, f"A{row}", rr, BLUE, PCT1, KEY_FILL, Alignment(horizontal="right"), BOX)
    for j in range(len(fm_list)):
        col = L(2 + j)
        put(wss, f"{col}{row}", f"={china_landed(f'$A{row}', f'{col}$18')}-$B$14", BLACK, USD3, border=BOX)
wss.conditional_formatting.add("B19:G25", FormulaRule(formula=["B19>0"], fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")))
wss.conditional_formatting.add("B19:G25", FormulaRule(formula=["B19<=0"], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))

put(wss, "A26", "Break-even BCD on China imports at base freight (India local = China)", BOLD, border=BOX)
put(wss, "B26", "=(((B14-B8-B9)/(B5+B6+B7)-1)-B11-B12)/(1+B10)", BOLD, PCT1, KEY_FILL, border=BOX)
put(wss, "A27", "Break-even freight multiplier at base BCD (India local = China)", BOLD, border=BOX)
put(wss, "B27", "=((B14-B8-B9)/(1+((1+B10)*B13+B11+B12))-(B5+B6))/B7", BOLD, MULT, KEY_FILL, border=BOX)
put(wss, "A28", "Check: grid formula at base BCD and 1.0x freight minus Calc (should be 0)", SUB)
put(wss, "B28", f"=ROUND({china_landed('$B$13', '1')}-$B$15,6)", SUB, '0.000000')

# Grid B: India advantage vs yield and ECMS pass-through
section(wss, 30, "Grid B: India advantage (USD/unit) vs India first-pass yield (rows) and ECMS pass-through to buyer (columns)", 8)
put(wss, "A31", "India reference values (base case)", BOLD)
irefs = [
    (32, "India: net material + conversion", f"{calc_pick(41, base_key(ind))}+{calc_pick(45, base_key(ind))}", USD3),
    (33, "India: fixed cost per unit", calc_pick(48, base_key(ind)), USD3),
    (34, "India: supplier margin", calc_pick(24, base_key(ind)), PCT1),
    (35, "India: ECMS incentive rate", calc_pick(35, base_key(ind)), PCT1),
    (36, "India: insurance + carrying factor (on price)", f"{calc_pick(30, base_key(ind))}+Inputs!$B$11*({calc_pick(27, base_key(ind))}*{calc_pick(29, base_key(ind))}+(1-{calc_pick(27, base_key(ind))})*{calc_pick(28, base_key(ind))})/365", '0.0000%'),
    (37, "India: international freight (trucking)", calc_pick(56, base_key(ind)), USD3),
]
for row, lab, fml, fmt in irefs:
    put(wss, f"A{row}", lab, BLACK, border=BOX)
    put(wss, f"B{row}", f"={fml}", GREEN, fmt, border=BOX)
put(wss, "A39", "India yield (rows) / ECMS pass-through (columns)", BOLD)
pt_list = [0.0, 0.25, 0.50, 0.75, 1.0]
y_list = [0.80, 0.83, 0.86, 0.89, 0.92, 0.95]
for j, p in enumerate(pt_list):
    put(wss, f"{L(2 + j)}39", p, BLUE, PCT1, KEY_FILL, CENTER, BOX)


def india_landed(y, pt):
    return f"((($B$32/{y}+$B$33)*(1+$B$34))*(1-$B$35*MIN(1,{pt}))*(1+$B$36)+$B$37)"


for i, y in enumerate(y_list):
    row = 40 + i
    put(wss, f"A{row}", y, BLUE, PCT1, KEY_FILL, Alignment(horizontal="right"), BOX)
    for j in range(len(pt_list)):
        col = L(2 + j)
        put(wss, f"{col}{row}", f"=$B$15-{india_landed(f'$A{row}', f'{col}$39')}", BLACK, USD3, border=BOX)
wss.conditional_formatting.add("B40:F45", FormulaRule(formula=["B40>0"], fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")))
wss.conditional_formatting.add("B40:F45", FormulaRule(formula=["B40<=0"], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
base_y = calc_pick(21, base_key(ind))
base_pt = calc_pick(36, base_key(ind))
put(wss, "A47", "Check: grid formula at base yield and pass-through minus Calc India landed (should be 0)", SUB)
put(wss, "B47", f"=ROUND({india_landed(base_y, base_pt)}-$B$14,6)", SUB, '0.000000')
put(wss, "A48", "Break-even India yield at base ECMS pass-through (India local = China)", BOLD, border=BOX)
put(wss, "B48", "=($B$32/((($B$15-$B$37)/((1+$B$34)*(1-$B$35*MIN(1," + base_pt + "))*(1+$B$36))-$B$33)))", BOLD, PCT1, KEY_FILL, border=BOX)
put(wss, "A49", "Check: India landed at the break-even yield minus China landed (should be 0)", SUB)
put(wss, "B49", "=ROUND(" + india_landed("$B$48", base_pt) + "-$B$15,6)", SUB, '0.000000')

# update Summary headline to include break-even yield
put(wsu, "A26", "Break-even India first-pass yield at which India local = China", BLACK, border=BOX)
put(wsu, "B26", "=Sensitivity!B48", BOLD, PCT1, border=BOX, align=Alignment(horizontal="right"))

# ---------------------------------------------------------------- INCOTERMS
wt = wb.create_sheet("Incoterms")
wt.sheet_view.showGridLines = False
wt.column_dimensions["A"].width = 46
for col in "BCDEFG":
    wt.column_dimensions[col].width = 15
wt.column_dimensions["H"].width = 3
put(wt, "A1", "Incoterms comparison: who pays what for an imported enclosure (USD per unit, base case)", TITLE)
put(wt, "A2", "Pick the exporting country in the yellow cell. FOB = buyer arranges carriage. CIF = seller pays freight and insurance to the Indian port. DDP = seller delivers duty paid to the plant.", SUB)
put(wt, "A3", "Exporting country", BOLD)
put(wt, "B3", "China", BLUE, fill=KEY_FILL, align=CENTER, border=BOX)
dv = DataValidation(type="list", formula1='"China,Vietnam"', allow_blank=False)
wt.add_data_validation(dv)
dv.add("B3")
key = f'Summary!$A$6&"|"&$B$3'
section(wt, 5, "Cost elements and who bears them", 7)
header_row(wt, 6, ["Cost element", "FOB: Seller", "FOB: Buyer", "CIF: Seller", "CIF: Buyer", "DDP: Seller", "DDP: Buyer"])
wt.row_dimensions[6].height = 30
elems = [
    (7, "Ex-works price (net of any incentive)", f"={calc_pick(53, key)}", "S", "S", "S"),
    (8, "International freight", f"={calc_pick(56, key)}", "B", "S", "S"),
    (9, "Cargo insurance", f"={calc_pick(57, key)}", "B", "S", "S"),
    (10, "Import duties (BCD + SWS + other)", f"={calc_pick(60, key)}+{calc_pick(61, key)}+{calc_pick(62, key)}", "B", "B", "S"),
    (11, "Customs clearance and port-to-plant haulage", f"={calc_pick(63, key)}+{calc_pick(64, key)}", "B", "B", "S"),
    (12, "Pipeline carrying cost (inventory financing)", f"={calc_pick(65, key)}", "B", "B", "S"),
]
for row, lab, fml, fob, cif, ddp in elems:
    put(wt, f"A{row}", lab, BLACK, border=BOX)
    put(wt, f"J{row}", fml, GREEN, USD3)  # helper value column
    for col_s, col_b, who in (("B", "C", fob), ("D", "E", cif), ("F", "G", ddp)):
        put(wt, f"{col_s}{row}", f"=IF(\"{who}\"=\"S\",$J{row},0)", BLACK, USD3, border=BOX)
        put(wt, f"{col_b}{row}", f"=IF(\"{who}\"=\"B\",$J{row},0)", BLACK, USD3, border=BOX)
wt.column_dimensions["J"].width = 14
put(wt, "J6", "Base value", WHITE_B, fill=HDR_FILL, align=CENTER, border=BOX)
put(wt, "A13", "Seller subtotal / Buyer subtotal (before markups)", BOLD, fill=TOT_FILL, border=BOX)
for col in "BCDEFG":
    put(wt, f"{col}13", f"=SUM({col}7:{col}12)", BOLD, USD3, TOT_FILL, border=BOX)
put(wt, "A14", "Seller service markup passed to buyer", BLACK, border=BOX)
put(wt, "B14", 0, BLACK, USD3, border=BOX)
put(wt, "C14", 0, BLACK, USD3, border=BOX)
put(wt, "D14", "", BLACK, border=BOX)
put(wt, "E14", "=Inputs!$B$15*($J$8+$J$9)", BLACK, USD3, border=BOX)
put(wt, "F14", "", BLACK, border=BOX)
put(wt, "G14", "=Inputs!$B$15*($J$8+$J$9)+Inputs!$B$16*($J$10+$J$11+$J$12)", BLACK, USD3, border=BOX)
put(wt, "A16", "Total cost to buyer (USD/unit)", BOLD, fill=TOT_FILL, border=BOX)
put(wt, "C16", "=SUM($J$7:$J$12)+C14", BOLD, USD3, TOT_FILL, border=BOX)
put(wt, "E16", "=SUM($J$7:$J$12)+E14", BOLD, USD3, TOT_FILL, border=BOX)
put(wt, "G16", "=SUM($J$7:$J$12)+G14", BOLD, USD3, TOT_FILL, border=BOX)
put(wt, "B16", "FOB", BOLD, fill=TOT_FILL, align=CENTER, border=BOX)
put(wt, "D16", "CIF", BOLD, fill=TOT_FILL, align=CENTER, border=BOX)
put(wt, "F16", "DDP", BOLD, fill=TOT_FILL, align=CENTER, border=BOX)
put(wt, "A17", "Premium vs FOB (%)", BLACK, border=BOX)
put(wt, "C17", "=C16/$C$16-1", BLACK, PCT1, border=BOX)
put(wt, "E17", "=E16/$C$16-1", BLACK, PCT1, border=BOX)
put(wt, "G17", "=G16/$C$16-1", BLACK, PCT1, border=BOX)
put(wt, "A18", "Check: FOB total equals Calc landed cost (should be 0)", SUB)
put(wt, "C18", f"=ROUND(C16-{calc_pick(66, key)},6)", SUB, '0.000000')

section(wt, 20, "Risk and control (qualitative)", 7)
qual = [
    ("Risk transfers to buyer at", "Loading on vessel/aircraft at origin", "Loading on vessel/aircraft at origin", "Delivery at the plant"),
    ("Who books and controls carriage", "Buyer", "Seller", "Seller"),
    ("Who bears Indian duty-rate change risk", "Buyer", "Buyer", "Seller (until re-priced)"),
    ("Who handles Indian customs clearance", "Buyer", "Buyer", "Seller"),
    ("Best fit when", "Buyer has scale freight contracts and customs expertise", "Seller is better placed to buy freight; buyer keeps duty control", "Buyer wants one landed price; accepts a premium for fewer touchpoints"),
]
header_row(wt, 21, ["", "FOB", "", "CIF", "", "DDP", ""])
for i, (lab, a, b_, c_) in enumerate(qual):
    row = 22 + i
    put(wt, f"A{row}", lab, BOLD, align=WRAP, border=BOX)
    for col_a, col_b, txt in (("B", "C", a), ("D", "E", b_), ("F", "G", c_)):
        wt.merge_cells(f"{col_a}{row}:{col_b}{row}")
        put(wt, f"{col_a}{row}", txt, BLACK, align=WRAP, border=BOX)
    wt.row_dimensions[row].height = 44 if i == 4 else 30

# ---------------------------------------------------------------- SOURCES
wso = wb.create_sheet("Sources")
wso.sheet_view.showGridLines = False
for col, w in zip("ABCD", [40, 22, 60, 80]):
    wso.column_dimensions[col].width = w
put(wso, "A1", "What is sourced and what is assumed", TITLE)
put(wso, "A2", "Retrieved 22 Oct 2026. Verify rates against the primary notification before relying on them.", SUB)
header_row(wso, 4, ["Input", "Status", "What the source says", "URL"])
src = [
    ("LME aluminium price", "Sourced", "LME cash offer USD 3,134.5/t on 6 Oct 2026", "https://www.alcircle.com/news/lme-aluminium-offer-price-rises-to-3-134-5-t-as-stocks-remain-unchanged-at-240-375t-121456"),
    ("ECMS incentive for enclosures", "Sourced (range)", "Electro-mechanicals and device enclosures get 6-8% turnover-linked incentive initially, 3-5% by end of scheme; 7% used as midpoint", "https://visionias.in/current-affairs/upsc-daily-news-summary/article/2025-04-09/the-hindu/schemes-in-news/govt-notifies-rs22919-crore-incentive-scheme-for-electronics-components"),
    ("ECMS covers mobile enclosures", "Sourced", "Cabinet-approved target segments include enclosures for mobile and IT hardware", "https://pmindia.gov.in/?p=16583068"),
    ("India BCD on metal phone parts", "Sourced (dated)", "BCD cut to 10% in Jan 2024 on mechanical items of plastic and metal; check the current CBIC schedule", "https://www.indiatvnews.com/topic/customs/news"),
    ("India supply-chain context", "Context", "Government urging deeper local component production (mechanicals, enclosures, batteries, chargers) before MacBook/iPad/Watch assembly", "https://www.storyboard18.com/brand-marketing/india-pushes-apple-suppliers-to-deepen-local-component-manufacturing-report-108791.htm"),
    ("India cost premium vs China", "Context (advocacy)", "Shareholder filing claims India/Vietnam production costs run 5-10% above China; treat as a claim, not a benchmark", "https://www.sec.gov/Archives/edgar/data/320193/000109690626000136/nlpc_px14a6g.htm"),
    ("US tariff regime (memo context)", "Context", "Section 301 forced-labor duty from 24 Jul 2026: India 10%, China and Vietnam 12.5%; China also keeps older Section 301 list duties", "https://gingercontrol.com/blog/importing-from-india-tariff-guide"),
    ("Cycle time, machine rate, labour rate, yield, tooling, overhead, margin, freight rates, transit days, insurance, clearance", "ASSUMPTION", "Illustrative values in plausible ranges. Replace with RFQ quotes, supplier cost breakdowns or benchmark studies before drawing real conclusions.", "n/a"),
]
for i, row_ in enumerate(src):
    r_ = 5 + i
    for j, v in enumerate(row_):
        put(wso, f"{L(1 + j)}{r_}", v, BOLD if (j == 1 and v == "ASSUMPTION") else BLACK, align=WRAP, border=BOX)
    wso.row_dimensions[r_].height = 48

# print setup: landscape, fit to one page wide
from openpyxl.worksheet.properties import PageSetupProperties
for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

# order sheets
order = ["README", "Summary", "Inputs", "Scenarios", "Sensitivity", "Incoterms", "Calc", "Sources"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("saved", OUT)
